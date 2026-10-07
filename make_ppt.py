#!/usr/bin/env python
"""make_ppt.py - build PowerPoint decks from documents, fully offline.

Commands:
  extract FILE...                 Dump text, tables and images to Markdown so an LLM can read them.
  build   SPEC.json [-o OUT.pptx] Render a deck from a JSON slide spec (see `schema`).
  auto    FILE... [-o OUT.pptx]   extract -> local Ollama model writes the spec -> build.
  schema                          Print the slide spec format.

Inputs: .pdf (scanned pages are OCR'd), .docx/.doc/.odt/.rtf, .xlsx/.xls/.ods/.csv,
.pptx/.ppt/.odp, images (.png/.jpg/...; OCR'd), .txt/.md.

Every build is rendered to PNG previews through LibreOffice so the result can be checked
without opening PowerPoint. Text is auto-shrunk to fit its box; anything that still does
not fit is reported as a warning.

Images (and pictures inside documents) are described by a local vision model when one is
installed (llama3.2-vision, llava, moondream...), so photos and diagrams get sensible slides.

Env overrides: MAKE_PPT_MODEL, OLLAMA_HOST, SOFFICE_PATH, TESSERACT_PATH.
"""
from __future__ import annotations

import argparse
import json
import logging
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import urllib.error
import urllib.request
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from families import FAMILIES, FamilyLayouts, tint

SOFFICE = os.environ.get("SOFFICE_PATH", r"C:\Program Files\LibreOffice\program\soffice.exe")
TESSERACT = os.environ.get("TESSERACT_PATH", r"C:\Program Files\Tesseract-OCR\tesseract.exe")
# Language data next to this script (eng, osd, hin). Falls back to Tesseract's own folder.
_local_tessdata = Path(__file__).with_name("tessdata")
TESSDATA = os.environ.get("TESSDATA_DIR") or (str(_local_tessdata) if _local_tessdata.is_dir() else "")
CS_FONT = "Nirmala UI"  # Windows font with Devanagari (and other Indic scripts)
OLLAMA = os.environ.get("OLLAMA_HOST", "http://localhost:11434").rstrip("/")
DEFAULT_MODEL = os.environ.get("MAKE_PPT_MODEL", "qwen3-coder:30b")

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp"}
LEGACY = {".doc": "docx", ".odt": "docx", ".rtf": "docx", ".xls": "xlsx", ".ods": "xlsx",
          ".ppt": "pptx", ".odp": "pptx"}
MAX_TABLE_ROWS = 40          # rows per sheet/table written to the extract
MIN_IMAGE_BYTES = 8_000      # skip icons and spacer images in documents

logging.getLogger("pdfminer").setLevel(logging.ERROR)  # FontBBox noise from pdfplumber
log = logging.getLogger("make_ppt")


# ---------------------------------------------------------------- cancel / skip
# The app sets these from the UI thread while a run is going; the CLI never touches them.

class Cancelled(Exception):
    """The user pressed Cancel."""


class Skipped(Exception):
    """The user skipped the stage that was running."""


STOP = threading.Event()   # cancel the whole run
SKIP: set[str] = set()     # stages to skip: "look" (describe pictures), "write" (stop retrying), "preview"


def checkpoint() -> None:
    if STOP.is_set():
        raise Cancelled("Cancelled.")


def interruptible(fn, stage: str | None = None):
    """Run a slow call (model reply, OCR, previews) on a helper thread so Cancel / Skip
    take effect at once instead of when the call returns."""
    box = {}

    def run():
        try:
            box["value"] = fn()
        except BaseException as e:  # SystemExit too: re-raised on the caller's thread
            box["error"] = e
    t = threading.Thread(target=run, daemon=True)
    t.start()
    while t.is_alive():
        t.join(0.25)
        checkpoint()
        if stage in SKIP:
            raise Skipped(stage)
    if "error" in box:
        raise box["error"]
    return box["value"]


# ---------------------------------------------------------------- extraction

def md_table(rows: list[list]) -> str:
    def cell(v) -> str:
        return "" if v is None else str(v).replace("|", "\\|").replace("\n", " ").strip()
    rows = [[cell(c) for c in r] for r in rows if r and any(c not in (None, "") for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def img_ref(p: Path) -> str:
    """Markdown image link relative to the work dir, which is where spec.json lives."""
    return f"![image](assets/{p.name})"


_langs: set[str] | None = None
OCR_USED: set[str] = set()     # languages actually used, recorded in the deck's properties
VISION_USED: set[str] = set()  # vision models actually used


def ocr_lang(img: Image.Image, config: str) -> str:
    """Pick the OCR language from the page's script: Devanagari -> Hindi + English."""
    import pytesseract
    global _langs
    if _langs is None:
        try:
            _langs = set(pytesseract.get_languages(config=config))
        except Exception:
            _langs = {"eng"}
    if "hin" not in _langs:
        lang = "eng"
    else:
        try:
            script = pytesseract.image_to_osd(img, config=config, output_type=pytesseract.Output.DICT)["script"]
            lang = "hin+eng" if script == "Devanagari" else "eng"
        except Exception:  # OSD needs enough text; hin+eng still reads English well
            lang = "hin+eng"
    OCR_USED.add(lang)
    return lang


def ocr(img: Image.Image) -> str:
    import pytesseract
    pytesseract.pytesseract.tesseract_cmd = TESSERACT
    config = f"--tessdata-dir {Path(TESSDATA).as_posix()}" if TESSDATA else ""
    try:
        return pytesseract.image_to_string(img, lang=ocr_lang(img, config), config=config).strip()
    except pytesseract.TesseractNotFoundError:
        log.warning("Tesseract not found at %s - skipping OCR", TESSERACT)
        return ""


def save_image(blob: bytes, assets: Path, name: str) -> Path | None:
    if len(blob) < MIN_IMAGE_BYTES:
        return None
    assets.mkdir(parents=True, exist_ok=True)
    out = assets / name
    out.write_bytes(blob)
    try:
        Image.open(out).verify()
    except Exception:
        out.unlink()  # not a format Pillow/PowerPoint can use (e.g. JBIG2, EMF)
        return None
    return out


def extract_pdf(path: Path, assets: Path) -> str:
    import pdfplumber
    import pypdf
    import pypdfium2
    parts, raster = [], None
    reader = pypdf.PdfReader(path)
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            checkpoint()
            text = (page.extract_text() or "").strip()
            label = f"## Page {i + 1}"
            scanned = len(text) < 20  # no text layer
            if scanned:
                raster = raster or pypdfium2.PdfDocument(path)
                text = ocr(raster[i].render(scale=2).to_pil())
                label += " (OCR)"
            parts.append(f"{label}\n\n{text}")
            for t in page.extract_tables():
                if t and max(map(len, t)) < 2:
                    continue  # one-column boxes are code/callouts, already in the text
                if table := md_table(t):
                    parts.append(f"Table on page {i + 1}:\n\n{table}")
            if scanned:
                continue  # its only "image" is the scan of the page itself
            try:
                for j, im in enumerate(reader.pages[i].images):
                    ext = Path(im.name).suffix or ".png"
                    if p := save_image(im.data, assets, f"p{i + 1}_img{j + 1}{ext}"):
                        parts.append(img_ref(p))
            except Exception as e:  # odd image encodings shouldn't kill the text extract
                log.warning("page %d: could not extract images (%s)", i + 1, e)
    if raster:
        raster.close()
    return "\n\n".join(parts)


def extract_docx(path: Path, assets: Path) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = docx.Document(path)
    parts = []
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.startswith("Heading") and style[-1:].isdigit():
                text = "#" * (int(style[-1]) + 1) + " " + text
            elif style == "Title":
                text = "# " + text
            elif "List" in style:
                text = "- " + text
            parts.append(text)
        elif el.tag == qn("w:tbl"):
            t = Table(el, doc)
            parts.append(md_table([[c.text for c in row.cells] for row in t.rows]))
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            name = Path(rel.target_part.partname).name
            if p := save_image(rel.target_part.blob, assets, name):
                parts.append(img_ref(p))
    return "\n\n".join(parts)


def extract_sheet(path: Path, assets: Path) -> str:
    import pandas as pd
    sheets = ({"data": pd.read_csv(path)} if path.suffix.lower() == ".csv"
              else pd.read_excel(path, sheet_name=None))
    parts = []
    for name, df in sheets.items():
        df = df.dropna(how="all").dropna(axis=1, how="all")
        parts.append(f"## Sheet: {name} ({len(df)} rows x {len(df.columns)} cols)")
        rows = [list(df.columns)] + df.head(MAX_TABLE_ROWS).values.tolist()
        parts.append(md_table(rows))
        if len(df) > MAX_TABLE_ROWS:
            parts.append(f"... {len(df) - MAX_TABLE_ROWS} more rows not shown.")
            numeric = df.select_dtypes("number")
            if not numeric.empty:
                parts.append("Column totals:\n\n" + md_table(
                    [["column", "sum", "mean"]] +
                    [[c, round(numeric[c].sum(), 2), round(numeric[c].mean(), 2)] for c in numeric]))
    return "\n\n".join(parts)


def extract_pptx(path: Path, assets: Path) -> str:
    parts = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        lines = [f"## Slide {i}"]
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            elif shape.has_table:
                lines.append(md_table([[c.text for c in r.cells] for r in shape.table.rows]))
            elif shape.shape_type == 13:  # picture
                img = shape.image
                if p := save_image(img.blob, assets, f"s{i}_{shape.shape_id}.{img.ext}"):
                    lines.append(img_ref(p))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            lines.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
        parts.append("\n\n".join(lines))
    return "\n\n".join(parts)


def extract_image(path: Path, assets: Path) -> str:
    assets.mkdir(parents=True, exist_ok=True)
    copy = assets / path.name
    shutil.copy2(path, copy)
    with Image.open(path) as img:
        size = f"{img.width}x{img.height}"
        text = ocr(img.convert("RGB"))
    out = f"Image file, {size} px:\n\n{img_ref(copy)}"
    return out + (f"\n\nText found in image (OCR):\n\n{text}" if text else "")


def convert_legacy(path: Path, fmt: str, outdir: Path) -> Path:
    if not Path(SOFFICE).exists():
        sys.exit(f"LibreOffice not found at {SOFFICE}; needed to read {path.suffix} files.")
    subprocess.run([SOFFICE, "--headless", "--convert-to", fmt, "--outdir", str(outdir), str(path)],
                   check=True, capture_output=True, timeout=300)
    out = outdir / f"{path.stem}.{fmt}"
    if not out.exists():
        sys.exit(f"LibreOffice could not convert {path}")
    return out


def extract(path: Path, assets: Path) -> str:
    ext = path.suffix.lower()
    if ext in LEGACY:
        with tempfile.TemporaryDirectory() as tmp:
            return extract(convert_legacy(path, LEGACY[ext], Path(tmp)), assets)
    handlers = {".pdf": extract_pdf, ".docx": extract_docx, ".xlsx": extract_sheet,
                ".xlsm": extract_sheet, ".csv": extract_sheet, ".pptx": extract_pptx}
    if ext in handlers:
        return handlers[ext](path, assets)
    if ext in IMAGE_EXT:
        return extract_image(path, assets)
    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace")
    sys.exit(f"Unsupported file type: {path}")


VISION_PREFS = ("llama3.2-vision", "qwen2.5vl", "gemma3", "minicpm-v", "llava", "moondream")
MAX_DESCRIBED = 12  # images described per run; each takes a few seconds
IMG_REF = re.compile(r"!\[image\]\(assets/([^)]+)\)")


VISION_PROMPT = ("Describe this image in 2-3 factual sentences for someone building a slide deck: "
                 "what it shows, and any text, numbers or labels visible. Do not guess.")


NO_SIGHT = re.compile(
    r"\b(unable to|not able to|can ?not|can't|don't have the ability to) (view|see|process|analy[sz]e|access|interpret) "
    r"(the |any |this )?(image|picture|photo)|\bas a text-based\b|\bno image (was|is) (provided|attached)", re.I)


_CAPS: dict[str, list | None] = {}


def model_caps(name: str) -> list | None:
    """What an Ollama model can do ('vision', 'thinking', ...), or None if Ollama doesn't say."""
    if name not in _CAPS:
        try:
            req = urllib.request.Request(f"{OLLAMA}/api/show", json.dumps({"model": name}).encode(),
                                         {"Content-Type": "application/json"})
            with urllib.request.urlopen(req, timeout=5) as r:
                _CAPS[name] = json.loads(r.read()).get("capabilities")
        except (urllib.error.URLError, OSError, ValueError):
            return None
    return _CAPS[name]


def vision_capable(name: str) -> bool:
    """Ask Ollama whether a model accepts images (newer Ollama reports 'capabilities');
    fall back to well-known vision model names."""
    if (caps := model_caps(name)) is not None:
        return "vision" in caps
    base = name.split(":")[0].lower()
    return base in VISION_PREFS or "vision" in base or base.endswith("vl")


def vision_candidates(preferred: str | None) -> list[str]:
    """The chosen vision model first, then installed *local* vision models, best first.
    Falling back to local models (never to an online one) keeps pictures on this PC."""
    try:
        with urllib.request.urlopen(f"{OLLAMA}/api/tags", timeout=5) as r:
            names = [m["name"] for m in json.loads(r.read())["models"]]
    except (urllib.error.URLError, OSError, KeyError, ValueError):
        names = []
    names = [n for n in names if not is_online(n)]
    found = [n for pref in VISION_PREFS for n in names if n.split(":")[0] == pref]
    if preferred:
        found = [preferred] + [n for n in found if n != preferred]
    return found


def describe_image(model: str, path: Path) -> str:
    """Short factual description of one picture: Ollama (local or cloud) or an API provider."""
    import base64
    import io
    with Image.open(path) as im:
        im = im.convert("RGB")
        im.thumbnail((1024, 1024))  # smaller = much faster, still enough detail
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=85)
    b64 = base64.b64encode(buf.getvalue()).decode()
    if model.startswith("api:"):
        import providers
        text = providers.describe_image(model, b64, VISION_PROMPT, 400, 300)
    else:
        text = ollama_stream("/api/generate", {
            "model": model, "keep_alive": 0,  # free memory for the slide model
            "prompt": VISION_PROMPT, "images": [b64], **think_option(model),
            "options": {"temperature": 0.2, "num_ctx": 4096, "num_predict": 400}}, 600, "look")[0].strip()
    if NO_SIGHT.search(text[:300]):  # text-only models may answer instead of erroring
        raise ValueError("model can't see images")
    return text


def add_image_descriptions(md: str, assets: Path, preferred: str | None) -> str:
    """Append a vision-model description after each image reference, so a text-only
    slide model knows what photos and diagrams show."""
    names = list(dict.fromkeys(IMG_REF.findall(md)))[:MAX_DESCRIBED]
    if not names:
        return md
    models = vision_candidates(preferred)
    if not models:
        log.warning("no offline vision model installed (e.g. `ollama pull moondream`); "
                    "images without text will not be described")
        return md
    if is_online(models[0]):
        log.warning("image model %s is online - the pictures will be sent to it", model_label(models[0]))
    descriptions = {}
    for name in names:
        if "look" in SKIP:
            log.info("skipped describing the remaining pictures")
            break
        while models:
            try:
                log.info("describing %s with %s", name, models[0])
                descriptions[name] = interruptible(lambda: describe_image(models[0], assets / name), "look")
                VISION_USED.add(models[0])
                break
            except Skipped:
                break
            except (urllib.error.HTTPError, ValueError) as e:  # memory, no image support, bad key...
                reason = e.code if isinstance(e, urllib.error.HTTPError) else e
                log.warning("%s failed (%s); trying the next vision model", models[0], reason)
                models.pop(0)
            except (urllib.error.URLError, OSError) as e:
                log.warning("could not describe %s (%s)", name, e)
                break
    return IMG_REF.sub(lambda m: m.group(0) + (
        f"\n\nImage description (vision model): {descriptions[m.group(1)]}"
        if descriptions.get(m.group(1)) else ""), md)


def extract_all(files: list[Path], workdir: Path, vision: str | None = "auto") -> Path:
    workdir.mkdir(parents=True, exist_ok=True)
    if missing := [str(f) for f in files if not f.exists()]:
        sys.exit("File not found: " + ", ".join(missing))
    chunks = []
    for f in files:
        checkpoint()
        log.info("extracting %s", f.name)
        chunks.append(f"# Source: {f.name}\n\n{extract(f, workdir / 'assets')}")
    text = "\n\n".join(chunks)
    if vision:  # "auto" = best installed vision model, None = skip
        text = add_image_descriptions(text, workdir / "assets", None if vision == "auto" else vision)
    md = workdir / "extracted.md"
    md.write_text(text, encoding="utf-8")
    return md


# ---------------------------------------------------------------- spec

SCHEMA = """{
  "title": "Deck title",
  "palette": "violet | midnight | forest | coral | terracotta | ocean | charcoal | teal | berry | cherry",
  "slides": [
    {"type": "title",   "title": "...", "subtitle": "...", "footer": "optional small text"},
    {"type": "section", "title": "...", "subtitle": "optional"},
    {"type": "bullets", "title": "...", "bullets": ["Lead phrase: detail", "..."],
                        "image": "optional path", "callout": "optional highlighted tip"},
    {"type": "cards",   "title": "...", "cards": [{"title": "...", "text": "...", "icon": "1-2 chars, optional"}]},
    {"type": "steps",   "title": "...", "steps": [{"title": "...", "text": "..."}]},
    {"type": "stats",   "title": "...", "stats": [{"value": "42%", "label": "..."}], "text": "optional"},
    {"type": "code",    "title": "...", "code": "multi-line string", "lang": "bash", "note": "optional warning"},
    {"type": "table",   "title": "...", "header": ["A", "B"], "rows": [["1", "2"]]},
    {"type": "chart",   "title": "...", "chart": "column | bar | line | pie",
                        "categories": ["Jan", "Feb"], "series": [{"name": "Sales", "values": [10, 20]}],
                        "text": "optional commentary shown beside the chart"},
    {"type": "image",   "title": "...", "image": "path", "caption": "optional", "bullets": ["optional"]},
    {"type": "closing", "title": "...", "items": ["checklist or key takeaways"]}
  ]
}
Every slide may also have "kicker" (small label above the title) and "notes" (speaker notes).
Top level may set "animate": false to turn off fade transitions and entrance animations."""

REQUIRED = {"title": ["title"], "section": ["title"], "bullets": ["title", "bullets"],
            "cards": ["title", "cards"], "steps": ["title", "steps"], "stats": ["title", "stats"],
            "code": ["title", "code"], "table": ["title", "header", "rows"],
            "chart": ["title", "categories", "series"], "image": ["title", "image"],
            "closing": ["title"]}


def validate(spec: dict, base: Path) -> list[str]:
    if not isinstance(spec, dict) or not isinstance(spec.get("slides"), list) or not spec["slides"]:
        return ['top level must be an object with a non-empty "slides" list']
    normalize(spec)
    errors = []
    for i, s in enumerate(spec["slides"], 1):
        t = s.get("type") if isinstance(s, dict) else None
        if t not in REQUIRED:
            errors.append(f"slide {i}: unknown type {t!r}; use one of {sorted(REQUIRED)}")
            continue
        errors += [f"slide {i} ({t}): missing {k!r}" for k in REQUIRED[t] if not s.get(k)]
        if t == "chart":
            n = len(s.get("categories") or [])
            for ser in s.get("series") or []:
                vals = ser.get("values") or []
                if len(vals) != n or not all(isinstance(v, (int, float)) for v in vals):
                    errors.append(f"slide {i}: series {ser.get('name')!r} needs {n} numeric values")
        img = s.get("image")
        if img and not isinstance(img, str):
            errors.append(f'slide {i}: "image" must be a path string like "assets/name.png"')
        elif img and not (base / img).exists():
            errors.append(f"slide {i}: image not found: {img}")
    return errors


def _text(v) -> str:
    """A list item a model wrote as an object -> plain text ('Lead: detail')."""
    if isinstance(v, dict):
        head = next((v[k] for k in ("title", "label", "heading", "name") if isinstance(v.get(k), str)), "")
        body = next((v[k] for k in ("text", "detail", "description", "content", "value") if isinstance(v.get(k), str)), "")
        return f"{head}: {body}" if head and body else head or body or json.dumps(v, ensure_ascii=False)
    return str(v)


DEFAULT_TITLES = {"title": "Overview", "section": "Next", "bullets": "Key points", "cards": "Highlights",
                  "steps": "Steps", "stats": "By the numbers", "code": "Command", "table": "Details",
                  "chart": "The numbers", "image": "Picture", "closing": "Key takeaways"}


def normalize(spec: dict) -> None:
    """Fix the shapes small local models commonly get wrong, in place, before validation:
    image as an object, a single bullet as a string, list items as objects, numbers as text."""
    for s in spec.get("slides", []) if isinstance(spec, dict) else []:
        if not isinstance(s, dict):
            continue
        if not str(s.get("title") or "").strip() and s.get("type") in DEFAULT_TITLES:
            # small models often drop titles; a placeholder beats failing the whole run
            s["title"] = s.get("kicker") or (spec.get("title") if s["type"] == "title" else "") \
                or DEFAULT_TITLES[s["type"]]
        img = s.get("image")
        if isinstance(img, list):
            img = img[0] if img else ""
        if isinstance(img, dict):
            s.setdefault("caption", img.get("caption") or img.get("alt") or "")
            img = next((img[k] for k in ("path", "src", "file", "url", "image") if isinstance(img.get(k), str)), "")
        if img:
            s["image"] = img
        else:
            s.pop("image", None)  # null / "" / unusable: drop it (required only on image slides)
        for key in ("bullets", "items"):
            if isinstance(s.get(key), str):
                s[key] = [s[key]]
            if isinstance(s.get(key), list):
                s[key] = [_text(x) for x in s[key]]
        for key, field in (("cards", "title"), ("steps", "title"), ("stats", "value")):
            if isinstance(s.get(key), list):  # items written as plain strings
                s[key] = [x if isinstance(x, dict) else {field: str(x)} for x in s[key]]
        for key in ("header", "categories"):
            if isinstance(s.get(key), list):
                s[key] = [_text(x) for x in s[key]]
        if isinstance(s.get("rows"), list):
            s["rows"] = [[_text(c) for c in r] if isinstance(r, list) else [_text(r)] for r in s["rows"]]
        for ser in s.get("series") or []:
            if isinstance(ser, dict) and isinstance(ser.get("values"), list):
                vals = []
                for v in ser["values"]:
                    try:
                        vals.append(float(str(v).replace(",", "").rstrip("%")) if isinstance(v, str) else v)
                    except ValueError:
                        vals.append(v)  # left for validate() to report
                ser["values"] = [int(v) if isinstance(v, float) and v.is_integer() else v for v in vals]


# ---------------------------------------------------------------- rendering

PALETTES = {  # dark (title/code bg), accent (dominant), accent2 (support)
    "violet": ("161A24", "6E4BF2", "149E8E"), "midnight": ("1E2761", "3D5AFE", "00897B"),
    "forest": ("1F3A2B", "2C7A3F", "B7791F"), "coral": ("2F3C7E", "E8505B", "D69E2E"),
    "terracotta": ("3A2A26", "B85042", "5F8A72"), "ocean": ("21295C", "065A82", "1C7293"),
    "charcoal": ("212121", "36454F", "D35400"), "teal": ("0B2E33", "028090", "02A37F"),
    "berry": ("2E1520", "6D2E46", "A26769"), "cherry": ("1C1F3B", "990011", "2F3C7E"),
}
WHITE, BODY, MUTED = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x2A, 0x2F, 0x3A), RGBColor(0x6B, 0x72, 0x80)
AMBER, AMBER_TINT = RGBColor(0xE8, 0x9A, 0x1C), RGBColor(0xFD, 0xF3, 0xE2)
CODE_FG, SOFT = RGBColor(0xE6, 0xE8, 0xEE), RGBColor(0xA0, 0xA6, 0xB4)
SANS, MONO = "Calibri", "Courier New"
W, H = 13.333, 7.5
LEFT, RIGHT, TOP, BOTTOM = 0.6, 12.733, 1.65, 6.9   # content area
CW = RIGHT - LEFT


def set_cs_font(run) -> None:
    """Complex-script font for the run, so Hindi (Devanagari) shows in a font that has it."""
    rPr = run._r.get_or_add_rPr()
    for old in rPr.findall(qn("a:cs")):
        rPr.remove(old)
    rPr.append(rPr.makeelement(qn("a:cs"), {"typeface": CS_FONT}))


def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr)


def lum(c: RGBColor) -> float:
    r, g, b = (v / 255 for v in c)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def mix(c: RGBColor, amount: float) -> RGBColor:
    """Blend colour towards white; amount=0.9 gives a pale tint."""
    return RGBColor(*(round(v + (255 - v) * amount) for v in c))


def text_height(paras: list[str], w: float, size: float, mono: bool, gap: float) -> float:
    cpl = max(1, int(w * 72 / (size * (0.6 if mono else 0.5))))  # chars per line, conservative
    lines = sum(max(1, math.ceil(len(p) / cpl)) for p in paras)
    return lines * size * 1.2 / 72 + len(paras) * gap / 72


class Deck(FamilyLayouts):
    def __init__(self, spec: dict, base: Path):
        self.spec, self.base, self.warnings = spec, base, []
        dark, accent, accent2 = PALETTES.get(spec.get("palette", "violet"), PALETTES["violet"])
        self.INK, self.ACCENT, self.ACCENT2 = rgb(dark), rgb(accent), rgb(accent2)
        self.ACCENT3, self.TINT = mix(self.ACCENT, 0.45), mix(self.ACCENT, 0.9)
        self.BG, self.BODY, self.HEAD, self.MUTED, self.soft = None, BODY, self.INK, MUTED, False
        self.ACCENTS = [self.ACCENT, self.ACCENT2, self.ACCENT3, mix(self.ACCENT2, 0.4)]
        self.fam = FAMILIES.get(spec.get("family") or "", {})
        self.HFONT = self.BFONT = self.NFONT = SANS
        self.caps, self.radius = False, None
        if f := self.fam:  # a design family: palette, fonts, card style and layout variants
            self.ACCENTS = [rgb(a) for a in f["accents"]]
            self.INK, (self.ACCENT, self.ACCENT2, self.ACCENT3) = rgb(f["dark"]), self.ACCENTS[:3]
            self.TINT, self.BG, self.BODY, self.MUTED = rgb(f["card"]), rgb(f["bg"]), rgb(f["ink"]), rgb(f["muted"])
            self.HEAD = self.BODY
            self.soft, self.caps, self.radius = f["soft"], f["caps"], f["radius"]
            self.HFONT, self.BFONT, self.NFONT = f["head"], f["body"], f["num"]
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.slide_no = self.section_no = 0
        self.anim = []

    # primitives ---------------------------------------------------------
    def text(self, slide, x, y, w, h, paras, size=16, color=None, bold=False, mono=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, min_size=10, gap=4, bullets=False, font=None):
        """paras: str or list; each item is str or list of (text, {bold, color}) runs."""
        paras = [paras] if isinstance(paras, (str, tuple)) else list(paras)
        plain = ["".join(r[0] for r in p) if isinstance(p, list) else str(p) for p in paras]
        inner_w = w - (0.3 if bullets else 0)
        fit = size
        while fit > min_size and text_height(plain, inner_w, fit, mono, gap) > h:
            fit -= 0.5
        if text_height(plain, inner_w, fit, mono, gap) > h * 1.05:
            self.warnings.append(f"slide {self.slide_no}: text may overflow: {plain[0][:50]!r}")
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(gap)
            for chunk, opts in (para if isinstance(para, list) else [(str(para), {})]):
                r = p.add_run()
                r.text = chunk
                f = r.font
                f.name, f.size = MONO if mono else (font or self.BFONT), Pt(fit)
                set_cs_font(r)
                f.bold = opts.get("bold", bold)
                f.color.rgb = opts.get("color", color or self.BODY)
            if bullets:
                pPr = p._p.get_or_add_pPr()
                pPr.set("marL", str(Inches(0.3)))
                pPr.set("indent", str(-Inches(0.3)))
                clr = pPr.makeelement(qn("a:buClr"), {})
                clr.append(clr.makeelement(qn("a:srgbClr"), {"val": str(self.ACCENT)}))
                pPr.append(clr)
                pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "\u25cf"}))
        return tb

    def box(self, slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=None, raised=True):
        s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = fill
        s.line.fill.background()
        s.shadow.inherit = False
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            fam_r = self.radius if self.radius is not None else 0.08
            s.adjustments[0] = fam_r if radius is None else (0 if fam_r == 0 else radius)
        if self.soft and raised and shape != MSO_SHAPE.RECTANGLE:  # soft shadow under light shapes
            eff = s._element.spPr.find(qn("a:effectLst"))
            shadow = eff.makeelement(qn("a:outerShdw"), {"blurRad": "203200", "dist": "50800",
                                                         "dir": "5400000", "algn": "t", "rotWithShape": "0"})
            clr = shadow.makeelement(qn("a:srgbClr"), {"val": "1A2333"})
            clr.append(clr.makeelement(qn("a:alpha"), {"val": "20000"}))
            shadow.append(clr)
            eff.append(shadow)
        return s

    def badge(self, slide, x, y, label, fill=None, d=0.6, size=18):
        c = self.box(slide, x, y, d, d, fill or self.ACCENT, MSO_SHAPE.OVAL, raised=False)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(label)[:2]
        light = lum(fill or self.ACCENT) > 0.62
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = self.NFONT, Pt(size), True, BODY if light else WHITE

    def code_block(self, slide, x, y, w, h, code, lang=""):
        self.box(slide, x, y, w, h, self.INK, radius=0.04)
        top = 0.2
        if lang:
            self.text(slide, x + 0.3, y + 0.18, w - 0.6, 0.3, lang.upper(), size=10,
                      color=mix(self.ACCENT2, 0.3), bold=True)
            top = 0.5
        self.text(slide, x + 0.3, y + top, w - 0.6, h - top - 0.15, code.rstrip().split("\n"),
                  size=14, color=CODE_FG, mono=True, min_size=9, gap=1)

    def warn(self, slide, x, y, w, h, msg):
        self.box(slide, x, y, w, h, AMBER_TINT)
        self.badge(slide, x + 0.25, y + (h - 0.5) / 2, "!", AMBER, d=0.5, size=20)
        self.text(slide, x + 0.95, y + 0.1, w - 1.15, h - 0.2, [lead(msg)], size=15, color=BODY,
                  anchor=MSO_ANCHOR.MIDDLE, min_size=11)

    def picture(self, slide, path, x, y, w, h):
        path = self.base / path
        with Image.open(path) as im:
            iw, ih = im.size
        scale = min(w / iw, h / ih)
        pw, ph = iw * scale, ih * scale
        slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                                 Inches(pw), Inches(ph))

    def new_slide(self, s, dark=False):
        self.slide_no += 1
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        if dark or self.BG:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = self.INK if dark else self.BG
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
        self.anim.append([slide, [0], False])  # [slide, group start indexes, skip first group]
        return slide

    def mark(self, slide):
        """Start a new animation group: shapes added after this fade in together."""
        self.anim[-1][1].append(len(slide.shapes))

    def header(self, s, title=None):
        slide = self.new_slide(s)
        style, tx, tw = self.fam.get("header"), LEFT, CW
        if style == "dots_page":  # three dots + raised page-number circle
            for i in range(3):
                self.box(slide, LEFT + i * 0.24, 0.3, 0.13, 0.13, self.ACCENTS[i], MSO_SHAPE.OVAL, raised=False)
            self.box(slide, RIGHT - 0.85, 0.35, 0.85, 0.85, self.TINT, MSO_SHAPE.OVAL)
            self.text(slide, RIGHT - 0.85, 0.35, 0.85, 0.85, f"{self.slide_no:02d}", size=20, bold=True,
                      font=self.NFONT, color=self.HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            tw = CW - 1.2
        elif style == "kicker_bar":  # accent bar beside the title
            self.box(slide, LEFT, 0.48, 0.09, 0.95, self.ACCENT, MSO_SHAPE.RECTANGLE, raised=False)
            tx, tw = LEFT + 0.32, CW - 0.32
        if s.get("kicker"):
            self.text(slide, tx, 0.35, tw, 0.3, s["kicker"].upper(), size=12, color=self.ACCENT, bold=True)
        self.text(slide, tx, 0.65, tw, 0.8, self.cap(title or s["title"]), size=34, color=self.HEAD,
                  bold=True, min_size=24, font=self.HFONT)
        if style == "stripes":
            for i in range(4):
                self.box(slide, LEFT + i * 0.45, 1.45, 0.35, 0.07, self.ACCENTS[i], MSO_SHAPE.RECTANGLE, raised=False)
        elif style == "hexdots":
            for i in range(5):
                self.box(slide, LEFT + i * 0.3, 1.43, 0.18, 0.16, self.ACCENTS[i % 4], MSO_SHAPE.HEXAGON, raised=False)
        elif style == "rule":
            self.box(slide, LEFT, 1.48, CW, 0.015, tint(self.MUTED, 0.55), MSO_SHAPE.RECTANGLE, raised=False)
        elif style == "kicker":
            self.box(slide, LEFT, 1.47, 0.7, 0.045, self.ACCENT, MSO_SHAPE.RECTANGLE, raised=False)
        self.mark(slide)
        self.anim[-1][2] = True  # title stays put; only content animates
        return slide

    # slide types --------------------------------------------------------
    def variant(self, kind, prefix, s):
        """Run the family's variant for this slide type, if it has one. True when handled."""
        fn = getattr(self, prefix + str(self.fam.get(kind)), None)
        if fn:
            fn(s)
        return fn is not None

    def s_title(self, s):
        if self.variant("title", "t_", s):
            return
        slide = self.new_slide(s, dark=True)
        self.text(slide, 0.8, 2.0, 11.7, 1.9, s["title"], size=54, color=WHITE, bold=True,
                  anchor=MSO_ANCHOR.BOTTOM, min_size=32)
        self.box(slide, 0.8, 4.15, 1.2, 0.12, self.ACCENT, MSO_SHAPE.RECTANGLE)
        self.mark(slide)
        if s.get("subtitle"):
            self.text(slide, 0.8, 4.5, 11, 1.4, s["subtitle"], size=22, color=CODE_FG, min_size=14)
        self.mark(slide)
        if s.get("footer"):
            self.text(slide, 0.8, 6.5, 11.5, 0.4, s["footer"], size=14, color=SOFT)

    def s_section(self, s):
        if self.variant("section", "sec_", s):
            return
        slide = self.new_slide(s, dark=True)
        self.section_no += 1
        self.badge(slide, 0.8, 2.3, s.get("icon") or str(self.section_no), self.ACCENT, d=0.9, size=26)
        self.text(slide, 0.8, 3.4, 11.7, 1.2, s["title"], size=44, color=WHITE, bold=True, min_size=28)
        self.mark(slide)
        if s.get("subtitle"):
            self.text(slide, 0.8, 4.6, 11, 1.2, s["subtitle"], size=20, color=CODE_FG, min_size=14)

    def s_bullets(self, s):
        items, has_side = s["bullets"], bool(s.get("image"))
        per = 5 if has_side else 6
        chunks = split_even(items, per)
        for n, chunk in enumerate(chunks):
            slide = self.header(s, s["title"] + (" (cont.)" if n else ""))
            bottom = BOTTOM - (1.15 if s.get("callout") and n == len(chunks) - 1 else 0)
            w = 6.6 if has_side else CW
            self.text(slide, LEFT, TOP + 0.1, w, bottom - TOP - 0.1, [lead(b) for b in chunk],
                      size=20, gap=14, min_size=12, bullets=True)
            self.mark(slide)
            if has_side:
                self.box(slide, 7.6, TOP, 5.13, bottom - TOP, self.TINT, radius=0.04)
                self.picture(slide, s["image"], 7.75, TOP + 0.15, 4.83, bottom - TOP - 0.3)
            self.mark(slide)
            if s.get("callout") and n == len(chunks) - 1:
                self.warn(slide, LEFT, BOTTOM - 0.95, CW, 0.95, s["callout"])

    def s_cards(self, s):
        if not self.variant("cards", "cd_", s):
            self.cards_default(s)

    def cards_default(self, s):
        cards = s["cards"][:6]
        cols = 2 if len(cards) == 4 else min(len(cards), 3)
        rows = math.ceil(len(cards) / cols)
        gap = 0.3
        cw = (CW - gap * (cols - 1)) / cols
        ch = min((BOTTOM - TOP - gap * (rows - 1)) / rows, 3.6 if rows == 1 else 2.2)
        slide = self.header(s)
        for i, c in enumerate(cards):
            self.mark(slide)
            x, y = LEFT + (i % cols) * (cw + gap), TOP + (i // cols) * (ch + gap)
            self.box(slide, x, y, cw, ch, self.TINT)
            icon_col = self.ACCENT if i % 2 == 0 else self.ACCENT2
            if rows == 1:  # tall cards: icon on top
                self.badge(slide, x + 0.3, y + 0.3, c.get("icon") or i + 1, icon_col, d=0.7, size=20)
                self.text(slide, x + 0.3, y + 1.2, cw - 0.6, 0.8, self.cap(c.get("title", "")), size=20,
                          color=self.HEAD, bold=True, min_size=14, font=self.HFONT)
                self.text(slide, x + 0.3, y + 2.05, cw - 0.6, ch - 2.3, c.get("text", ""), size=16, min_size=11)
            else:  # short cards: icon at left
                self.badge(slide, x + 0.3, y + 0.3, c.get("icon") or i + 1, icon_col, d=0.6)
                self.text(slide, x + 1.15, y + 0.25, cw - 1.45, 0.5, self.cap(c.get("title", "")), size=18,
                          color=self.HEAD, bold=True, min_size=13, font=self.HFONT)
                self.text(slide, x + 1.15, y + 0.8, cw - 1.45, ch - 1.0, c.get("text", ""), size=14, min_size=10)

    def s_steps(self, s):
        if self.variant("steps", "st_", s):
            return
        steps = s["steps"]
        chunks = split_even(steps, 5)
        num = 0
        for n, chunk in enumerate(chunks):
            slide = self.header(s, s["title"] + (" (cont.)" if n else ""))
            gap = 0.2
            rh = min(1.4, (BOTTOM - TOP - gap * (len(chunk) - 1)) / len(chunk))
            for i, st in enumerate(chunk):
                self.mark(slide)
                num += 1
                y = TOP + i * (rh + gap)
                self.box(slide, LEFT, y, CW, rh, self.TINT)
                self.badge(slide, LEFT + 0.3, y + (rh - 0.6) / 2, num)
                self.text(slide, LEFT + 1.2, y + 0.12, 3.6, rh - 0.24, self.cap(st.get("title", "")), size=18,
                          color=self.HEAD, bold=True, anchor=MSO_ANCHOR.MIDDLE, min_size=13, font=self.HFONT)
                self.text(slide, LEFT + 5.0, y + 0.12, CW - 5.3, rh - 0.24, st.get("text", ""), size=15,
                          anchor=MSO_ANCHOR.MIDDLE, min_size=10)

    def s_stats(self, s):
        if self.variant("stats", "sa_", s):
            return
        stats = s["stats"][:4]
        slide = self.header(s)
        gap = 0.3
        cw = (CW - gap * (len(stats) - 1)) / len(stats)
        ch = 2.6 if s.get("text") else 3.0  # a fixed, compact card; full height leaves it mostly empty
        for i, st in enumerate(stats):
            self.mark(slide)
            x = LEFT + i * (cw + gap)
            self.box(slide, x, TOP, cw, ch, self.TINT)
            self.text(slide, x + 0.3, TOP + 0.3, cw - 0.6, 1.3, str(st.get("value", "")), size=60,
                      color=self.ACCENTS[i % len(self.ACCENTS)], bold=True,
                      anchor=MSO_ANCHOR.BOTTOM, min_size=28, font=self.NFONT)
            self.text(slide, x + 0.3, TOP + 1.7, cw - 0.6, ch - 1.9, st.get("label", ""), size=16, min_size=11)
        self.mark(slide)
        if s.get("text"):
            self.text(slide, LEFT, TOP + ch + 0.4, CW, BOTTOM - TOP - ch - 0.4, [lead(s["text"])],
                      size=18, min_size=12)

    def s_code(self, s):
        slide = self.header(s)
        y = TOP
        self.mark(slide)
        if s.get("note"):
            self.warn(slide, LEFT, y, CW, 0.95, s["note"])
            y += 1.2
        lines = s["code"].rstrip().split("\n")
        h = min(BOTTOM - y, 0.85 + text_height(lines, CW - 0.6, 14, True, 1))  # counts wrapped lines
        self.mark(slide)
        self.code_block(slide, LEFT, y, CW, max(h, 1.1), s["code"], s.get("lang", ""))

    def s_table(self, s):
        header, rows = s["header"], s["rows"]
        chunks = split_even(rows, 8) or [[]]
        lens = [max(len(str(r[c])) if c < len(r) else 0 for r in [header] + rows) for c in range(len(header))]
        weights = [min(max(l, 6), 40) for l in lens]
        size = 15 if len(header) <= 3 else 13 if len(header) <= 5 else 11
        for n, chunk in enumerate(chunks):
            slide = self.header(s, s["title"] + (" (cont.)" if n else ""))
            data = [header] + chunk
            row_h = min(0.6, (BOTTOM - TOP) / len(data))
            tbl = slide.shapes.add_table(len(data), len(header), Inches(LEFT), Inches(TOP),
                                         Inches(CW), Inches(row_h * len(data))).table
            for c, wt in enumerate(weights):
                tbl.columns[c].width = Emu(int(Inches(CW) * wt / sum(weights)))
            for r, row in enumerate(data):
                for c in range(len(header)):
                    cell = tbl.cell(r, c)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = self.INK if r == 0 else (self.TINT if r % 2 else (self.BG or WHITE))
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    cell.margin_left = cell.margin_right = Inches(0.15)
                    run = cell.text_frame.paragraphs[0].add_run()
                    run.text = str(row[c]) if c < len(row) else ""
                    run.font.size, run.font.name = Pt(size), SANS
                    set_cs_font(run)
                    run.font.bold = r == 0
                    run.font.color.rgb = WHITE if r == 0 else self.BODY

    def s_chart(self, s):
        slide = self.header(s)
        kind = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
                "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}.get(
                    s.get("chart", "column"), XL_CHART_TYPE.COLUMN_CLUSTERED)
        data = CategoryChartData()
        data.categories = [str(c) for c in s["categories"]]
        for ser in s["series"]:
            data.add_series(ser.get("name", ""), ser["values"])
        cw = 8.0 if s.get("text") else CW
        chart = slide.shapes.add_chart(kind, Inches(LEFT), Inches(TOP), Inches(cw),
                                       Inches(BOTTOM - TOP), data).chart
        chart.font.size, chart.font.name = Pt(12), SANS
        chart.font.color.rgb = self.BODY
        palette = [self.ACCENT, self.ACCENT2, self.ACCENT3, self.INK, AMBER, self.MUTED]
        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(11)
        pie = kind == XL_CHART_TYPE.PIE
        chart.has_legend = pie or len(s["series"]) > 1
        if chart.has_legend:
            chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
        if pie:
            for i, pt in enumerate(plot.series[0].points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = palette[i % len(palette)]
        else:
            if hasattr(plot, "gap_width"):
                plot.gap_width = 80
            for i, ser in enumerate(plot.series):
                col = palette[i % len(palette)]
                if kind == XL_CHART_TYPE.LINE_MARKERS:
                    ser.format.line.color.rgb = col
                    ser.format.line.width = Pt(2.5)
                else:
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = col
            chart.value_axis.has_major_gridlines = True
            chart.value_axis.major_gridlines.format.line.color.rgb = (
                mix(self.BG, 0.12) if self.BG and lum(self.BG) < 0.45 else RGBColor(0xE5, 0xE7, 0xEB))
            chart.value_axis.format.line.fill.background()
            if kind == XL_CHART_TYPE.BAR_CLUSTERED:  # PowerPoint draws bar categories bottom-up
                chart.category_axis.reverse_order = True
        self.mark(slide)
        if s.get("text"):
            self.box(slide, 8.9, TOP, RIGHT - 8.9, BOTTOM - TOP, self.TINT)
            self.text(slide, 9.2, TOP + 0.3, RIGHT - 9.5, BOTTOM - TOP - 0.6, [lead(s["text"])],
                      size=17, min_size=11)

    def s_image(self, s):
        slide = self.header(s)
        side = bool(s.get("bullets"))
        w = 7.6 if side else CW
        cap_h = 0.5 if s.get("caption") else 0
        self.box(slide, LEFT, TOP, w, BOTTOM - TOP - cap_h, self.TINT, radius=0.04)
        self.picture(slide, s["image"], LEFT + 0.2, TOP + 0.2, w - 0.4, BOTTOM - TOP - cap_h - 0.4)
        if s.get("caption"):
            self.text(slide, LEFT, BOTTOM - 0.4, w, 0.4, s["caption"], size=12, color=self.MUTED,
                      align=PP_ALIGN.CENTER)
        self.mark(slide)
        if side:
            self.text(slide, LEFT + w + 0.4, TOP + 0.1, CW - w - 0.4, BOTTOM - TOP - 0.1,
                      [lead(b) for b in s["bullets"]], size=18, gap=12, min_size=11, bullets=True)

    def s_closing(self, s):
        if self.variant("closing", "cl_", s):
            return
        slide = self.new_slide(s, dark=True)
        items = s.get("items", [])[:6]
        self.text(slide, 0.8, 0.7, 11.7, 0.9, s["title"], size=40, color=WHITE, bold=True, min_size=26)
        step = min(0.95, 4.9 / max(len(items), 1))
        for i, item in enumerate(items):
            self.mark(slide)
            y = 1.95 + i * step
            self.badge(slide, 0.8, y, "\u2713", self.ACCENT2, d=0.6, size=20)
            self.text(slide, 1.65, y, 10.8, 0.6, item, size=22, color=WHITE,
                      anchor=MSO_ANCHOR.MIDDLE, min_size=13)

    def build(self, out: Path, animate: bool = True) -> Path:
        self.slide_map = []  # spec slide -> its first deck slide (long slides continue onto more)
        for s in self.spec["slides"]:
            self.slide_map.append(len(self.prs.slides))
            getattr(self, "s_" + s["type"])(s)
        if animate and self.spec.get("animate", True):
            for slide, marks, skip in self.anim:
                shapes = list(slide.shapes)
                bounds = list(zip(marks, marks[1:] + [len(shapes)]))[1 if skip else 0:]
                add_animation(slide, [shapes[a:b] for a, b in bounds if b > a])
        self.prs.core_properties.title = self.spec.get("title", "")
        self.prs.core_properties.comments = self.spec.get(
            "made_with", "Built by make_ppt.py (offline) from a hand-written spec; no AI model involved.")[:255]
        out.parent.mkdir(parents=True, exist_ok=True)
        try:
            self.prs.save(out)
        except PermissionError:
            sys.exit(f"Cannot write {out} - is it open in PowerPoint?")
        return out


P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
FADE_MS, STAGGER_MS = 500, 300


def add_animation(slide, groups: list[list]) -> None:
    """Fade transition, plus each group fading in automatically, one after another.

    python-pptx has no animation API, so this writes the <p:transition> and <p:timing>
    XML that PowerPoint itself produces for an "After Previous" fade entrance.
    """
    from lxml import etree
    sld = slide._element
    for tag in ("transition", "timing"):
        for old in sld.findall(qn(f"p:{tag}")):
            sld.remove(old)
    anchor = sld.find(qn("p:clrMapOvr"))
    if anchor is None:
        anchor = sld.find(qn("p:cSld"))
    transition = etree.fromstring(
        f'<p:transition xmlns:p="{P_NS}" spd="med"><p:fade/></p:transition>')
    anchor.addnext(transition)
    if not groups:
        return
    ids = iter(range(4, 100_000))
    effects, builds = [], []
    for g, shapes in enumerate(groups):
        inner = []
        for shape in shapes:
            spid = shape.shape_id
            a, b, c = next(ids), next(ids), next(ids)
            node = "afterEffect" if g == 0 and not inner else "withEffect"
            tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
            inner.append(
                f'<p:par><p:cTn id="{a}" presetID="10" presetClass="entr" presetSubtype="0" '
                f'fill="hold" grpId="0" nodeType="{node}">'
                f'<p:stCondLst><p:cond delay="{g * STAGGER_MS}"/></p:stCondLst><p:childTnLst>'
                f'<p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/>'
                f'</p:stCondLst></p:cTn>{tgt}<p:attrNameLst><p:attrName>style.visibility</p:attrName>'
                f'</p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{c}" dur="{FADE_MS}"/>'
                f'{tgt}</p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par>')
            if shape.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX):
                builds.append(f'<p:bldP spid="{spid}" grpId="0" animBg="1"/>')
        effects += inner
    bld = f'<p:bldLst>{"".join(builds)}</p:bldLst>' if builds else ""
    timing = etree.fromstring(
        f'<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par>'
        f'<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
        f'<p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
        f'<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
        f'<p:par><p:cTn id="{next(ids)}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        f'<p:childTnLst>{"".join(effects)}</p:childTnLst></p:cTn></p:par>'
        f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
        f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
        f'</p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/>'
        f'</p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
        f'{bld}</p:timing>')
    transition.addnext(timing)


def split_even(items: list, per: int) -> list[list]:
    """Split into the fewest chunks of <= per items, balanced (7 -> 4+3, not 6+1)."""
    if not items:
        return []
    n = math.ceil(len(items) / per)
    size = math.ceil(len(items) / n)
    return [items[i:i + size] for i in range(0, len(items), size)]


def lead(s: str) -> list:
    """'Lead phrase: rest' -> bold lead run + normal run."""
    head, sep, rest = str(s).partition(": ")
    if sep and len(head) <= 45 and rest:
        return [(head + ": ", {"bold": True}), (rest, {})]
    return [(str(s), {})]


def render_previews(deck: Path, outdir: Path) -> list[Path]:
    """PPTX -> PDF (LibreOffice) -> one PNG per slide + a contact sheet."""
    import pypdfium2
    if not Path(SOFFICE).exists():
        log.warning("LibreOffice not found at %s - skipping previews", SOFFICE)
        return []
    # Reuse the folder: deleting and recreating it right away fails on Windows with
    # "Access is denied" (the delete is still pending, or Explorer has the folder open).
    outdir.mkdir(parents=True, exist_ok=True)
    for old in [*outdir.glob("slide-*.png"), outdir / "contact-sheet.png"]:
        old.unlink(missing_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "deck.pptx"
        shutil.copy2(deck, src)
        subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                       check=True, capture_output=True, timeout=300)
        pdf = pypdfium2.PdfDocument(Path(tmp) / "deck.pdf")
        pngs = []
        for i in range(len(pdf)):
            p = outdir / f"slide-{i + 1:02d}.png"
            pdf[i].render(scale=1.2).to_pil().save(p)
            pngs.append(p)
        pdf.close()
    thumbs = [Image.open(p) for p in pngs]
    tw, th = thumbs[0].width // 2, thumbs[0].height // 2
    sheet = Image.new("RGB", (tw * 2, th * math.ceil(len(thumbs) / 2)), "white")
    for i, im in enumerate(thumbs):
        sheet.paste(im.resize((tw, th)), ((i % 2) * tw, (i // 2) * th))
    sheet.save(outdir / "contact-sheet.png")
    return pngs


def apply_design(spec: dict, design: str | None) -> None:
    """Use a design family (see families.py / DESIGN_NOTES.md). None/"" keeps the default look."""
    if not design:
        return
    if design not in FAMILIES:
        sys.exit(f"Unknown design '{design}'. Choose one of: {', '.join(FAMILIES)}")
    spec["family"] = design
    spec.pop("theme", None)


def build(spec_path: Path, out: Path | None, preview: bool = True, animate: bool = True,
          design: str | None = None) -> Path:
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    if design:
        apply_design(spec, design)
        spec_path.write_text(json.dumps(spec, indent=2, ensure_ascii=False), encoding="utf-8")
    base = spec_path.parent
    if errors := validate(spec, base):
        sys.exit("Spec errors:\n  " + "\n  ".join(errors))
    out = out or base / f"{safe_name(spec.get('title') or spec_path.stem)}.pptx"
    deck = Deck(spec, base)
    deck.build(out, animate)
    (base / "slide_map.json").write_text(json.dumps(deck.slide_map), encoding="utf-8")
    print(f"Deck: {out} ({deck.slide_no} slides)")
    for w in deck.warnings:
        print(f"  warning: {w}")
    if preview and "preview" not in SKIP:
        try:
            pngs = interruptible(lambda: render_previews(out, base / "preview"), "preview")
        except Skipped:
            log.info("skipped the slide previews")
            pngs = []
        except (OSError, subprocess.SubprocessError) as e:  # previews are a convenience, the deck is done
            log.warning("deck saved, but previews could not be made (%s) - close the preview folder or "
                        "any open slide images and rebuild to get them", e)
            pngs = []
        if pngs:
            print(f"Previews: {pngs[0].parent} (contact-sheet.png shows all slides)")
    return out


def safe_name(s: str) -> str:
    return "".join(c for c in s if c not in '<>:"/\\|?*').strip()[:80] or "deck"


# ---------------------------------------------------------------- auto (Ollama)

PROMPT = """You turn source documents into a clear, well-structured slide deck.

Reply with ONLY a JSON object in this format:
{schema}

Rules:
- {count} Start with a "title" slide, end with a "closing" slide.
- Use at most one "section" slide per 5 slides, and none in decks under 10 slides.
- Vary slide types. Prefer cards, steps, stats, tables and charts over plain bullets.
- Keep the whole reply short: tables at most 8 rows, never copy a long table whole.
- Max 6 bullets per slide, max 20 words each. A bullet may start with a short bold label and a colon, e.g. "Budget: 12 lakh".
- Put the detail that does not fit on the slide into "notes".
- Use only facts from the source. Do not invent numbers.
- If the source contains commands or code, include them in "code" slides, copied exactly.
  Commands are usually the most important content of a technical guide - never drop them.
- Only make a "chart" when the source has a table of numbers. Never invent chart data.
- Only use image paths listed in the source as ![image](path), copied exactly.
- Use the "Image description" lines to write titles, captions and bullets for images.
  When the source is mostly images, give each important image its own slide.
- Write the slides in the same language as the source (Hindi source -> Hindi slides),
  unless the instructions ask for another language. Keep commands and numbers as they are.
- Pick the palette that best suits the topic.
{extra}
SOURCE DOCUMENTS:
{content}"""


RETIRED_FILE = Path(__file__).with_name(".retired_models")  # cloud models Ollama has retired


def retired_models() -> set[str]:
    try:
        return set(RETIRED_FILE.read_text(encoding="utf-8").split())
    except OSError:
        return set()


def mark_retired(model: str) -> None:
    """Remember a retired cloud model so menus stop offering it (its local tag still exists)."""
    try:
        RETIRED_FILE.write_text("\n".join(sorted(retired_models() | {model})), encoding="utf-8")
    except OSError:
        pass


def is_online(model: str) -> bool:
    """API-provider models and Ollama cloud models run on someone else's servers."""
    return model.startswith("api:") or model.endswith(":cloud") or "-cloud" in model


def model_label(model: str) -> str:
    if model.startswith("api:"):
        import providers
        return providers.label(model)
    return f"{model} (Ollama cloud)" if is_online(model) else f"{model} (Ollama, local)"


def ask_llm(model: str, prompt: str) -> str:
    """Route to an OpenAI-compatible API provider (api:<id>:<model>) or to Ollama."""
    if model.startswith("api:"):
        import providers
        try:
            return interruptible(lambda: providers.chat_json(model, prompt, MAX_REPLY_TOKENS, REPLY_TIMEOUT), "write")
        except ValueError as e:
            sys.exit(str(e))
    return interruptible(lambda: ask_ollama(model, prompt), "write")


MAX_REPLY_TOKENS = 6144   # a 20-slide spec is ~4k tokens
THINK_TOKENS = 6144       # extra room for a thinking model's reasoning
MAX_CTX = 32768           # biggest context window asked of a local model
# Worst case seen: qwen3 read a number-heavy price list as 1.83 chars/token (phi4: 2.5).
# Under-estimating is what breaks: Ollama then drops the start of the prompt - the schema and
# rules - and the model writes placeholder slides.
CHARS_PER_TOKEN = 1.8


def reply_budget(model: str) -> int:
    """num_predict for a model; thinking counts against it, so thinking models get more."""
    return MAX_REPLY_TOKENS + (THINK_TOKENS if not model.startswith("api:") and thinks(model) else 0)


def source_limit(model: str) -> int:
    """Characters of source that still leave room for the rules and the reply."""
    if is_online(model):
        return 60_000
    return int((MAX_CTX - reply_budget(model) - 2500) * CHARS_PER_TOKEN)
REPLY_TIMEOUT = 1800      # seconds per model reply


def thinks(model: str) -> bool:
    return "thinking" in (model_caps(model) or [])


def think_option(model: str) -> dict:
    """gpt-oss thinks at length by default and used its whole reply budget doing so: keep it low.
    qwen3 / deepseek-r1 keep thinking - with it off their specs ran away or came out as
    placeholders - and get extra reply budget instead (see ask_ollama)."""
    return {"think": "low"} if model.startswith("gpt-oss") and thinks(model) else {}


def ollama_stream(path: str, payload: dict, timeout: int, stage: str) -> tuple[str, str]:
    """POST to Ollama with streaming on; returns (text, done_reason). Streaming lets a
    cancelled or skipped run close the connection, which makes Ollama stop generating."""
    payload["stream"] = True
    req = urllib.request.Request(f"{OLLAMA}{path}", json.dumps(payload).encode(),
                                 {"Content-Type": "application/json"})
    text, reason = [], ""
    with urllib.request.urlopen(req, timeout=timeout) as r:
        for line in r:
            if STOP.is_set() or stage in SKIP:
                break  # leaving the `with` closes the connection
            chunk = json.loads(line)
            if chunk.get("error"):
                raise ValueError(chunk["error"])
            text.append(chunk.get("response") or (chunk.get("message") or {}).get("content") or "")
            reason = chunk.get("done_reason") or reason
    return "".join(text), reason


def ask_ollama(model: str, prompt: str) -> str:
    # num_predict caps the reply: in JSON mode small models can loop forever (e.g. endless
    # whitespace); a capped, cut-off reply just fails validation and gets retried instead of
    # hanging until the timeout.
    budget = reply_budget(model)
    # Context sized to the prompt + room for the reply, in 4k steps: a fixed 32k window
    # reserves several GB of memory even for a short document.
    num_ctx = min(MAX_CTX, math.ceil((len(prompt) / CHARS_PER_TOKEN + budget) / 4096) * 4096)
    think = think_option(model)
    payload = {"model": model, "format": "json", "messages": [{"role": "user", "content": prompt}], **think,
               "options": {"temperature": 0.3, "num_ctx": num_ctx, "num_predict": budget, "repeat_penalty": 1.1}}
    try:
        text, reason = ollama_stream("/api/chat", payload, REPLY_TIMEOUT, "write")
        if reason == "length" and not text.strip():
            log.warning("%s used its whole reply budget before writing anything", model)
        return text
    except TimeoutError:  # raised directly (not as URLError) when the reply is too slow
        sys.exit(f"{model} did not finish within {REPLY_TIMEOUT // 60} minutes. It may be too big for the "
                 "free memory (running partly on the CPU) - try a smaller model or fewer slides.")
    except urllib.error.HTTPError as e:  # Ollama is up but refused, e.g. not enough memory
        try:
            reason = json.loads(e.read()).get("error", str(e))
        except Exception:
            reason = str(e)
        if "retired" in reason.lower() or e.code == 410:
            mark_retired(model)
            sys.exit(f"{model} no longer exists: Ollama retired this cloud model ({reason}). It is now hidden "
                     f"from the app's menus. Pick another model; you can also remove it with `ollama rm {model}`.")
        hint = (" Close other heavy apps, or pick a smaller model with -m." if "memory" in reason else
                " Ollama cloud models need you to be signed in: run `ollama signin`."
                if "unauthorized" in reason.lower() or e.code == 401 else "")
        sys.exit(f"Ollama error: {reason}.{hint}")
    except urllib.error.URLError as e:
        sys.exit(f"Cannot reach Ollama at {OLLAMA} ({e}). Is `ollama serve` running?")


CODE_HINTS = ("$ ", "PS>", ".exe", "--", "sudo ", "pip ", "npm ", "curl ", "git ", "cd ", "./", ".\\")


def content_checks(spec: dict, source: str) -> list[str]:
    """Catch the two mistakes local models make most: invented charts and dropped commands."""
    errors = []
    numbers = set(re.findall(r"\d+(?:\.\d+)?", source.replace(",", "")))
    for i, s in enumerate(spec.get("slides", []), 1):
        if s.get("type") != "chart":
            continue
        vals = [v for ser in s.get("series") or [] for v in ser.get("values") or []]
        missing = [v for v in vals if f"{v:g}" not in numbers]
        if missing or len(set(vals)) < 2:
            errors.append(f"slide {i}: chart values {missing or vals} are not real data from the source; remove this chart or use a table/cards instead")
    code_lines = [l for l in source.splitlines()  # markdown table rows ("|---|") are not commands
                  if any(h in l for h in CODE_HINTS) and not l.lstrip().startswith("|")]
    if len(code_lines) >= 3 and not any(s.get("type") == "code" for s in spec.get("slides", [])):
        errors.append("the source contains commands (e.g. " + repr(code_lines[0].strip()[:60]) + ") but there is no "
                      "\"code\" slide; add code slides with the commands copied exactly")
    slides = spec.get("slides", [])
    n_sec = sum(s.get("type") == "section" for s in slides)
    if n_sec > max(len(slides) // 5, 0 if len(slides) < 10 else 1):
        errors.append(f"{n_sec} section slides is too many for {len(slides)} slides; remove section slides and keep the content slides")
    return errors


CHART_WORDS = re.compile(r"\b(chart|graph|plot|bar|axis|axes)\b", re.I)


def add_missing_images(spec: dict, files: list[Path], source: str) -> list[str]:
    """Make sure every picture the user gave ends up in the deck.

    Small models often skip some. Each skipped input image gets its own image slide
    (before the closing slide), captioned from its vision description. A chart picture
    is not re-added when the model already rebuilt it as a native chart slide.
    """
    slides = spec["slides"]
    used = {Path(s.get("image") or "").name for s in slides}
    has_chart = any(s.get("type") == "chart" for s in slides)
    descriptions = dict(re.findall(
        r"!\[image\]\(assets/([^)]+)\)\s*\n\nImage description \(vision model\): ([^\n]+)", source))
    added = []
    for f in files:
        if f.suffix.lower() not in IMAGE_EXT or f.name in used:
            continue
        desc = descriptions.get(f.name, "")
        if has_chart and CHART_WORDS.search(desc):
            continue
        first = re.split(r"(?<=[.!?])\s", desc, maxsplit=1)[0] if desc else ""
        first = re.sub(r"^(The|This) image (depicts|shows|presents|is|features)\s+", "", first)
        first = first[:1].upper() + first[1:]
        slide = {"type": "image", "title": f.stem.replace("_", " ").replace("-", " ").title(),
                 "image": f"assets/{f.name}", "notes": desc}
        if first:
            slide["caption"] = first[:150]
        at = len(slides) - 1 if slides and slides[-1].get("type") == "closing" else len(slides)
        slides.insert(at, slide)
        added.append(f.name)
    return added


SLIDE_PROMPT = """You edit one slide of an existing slide deck.

Reply with ONLY one JSON object: a single slide in one of the formats inside "slides" here:
{schema}

The deck "{title}" has these slides:
{outline}

TASK: {task}
{instruction}
Rules: max 6 bullets, max 20 words each; use only facts from the source; do not invent numbers;
write in the same language as the other slides unless told otherwise; image paths only as listed
in the source as ![image](path).

SOURCE DOCUMENTS:
{content}"""


def ai_slide(model: str, spec: dict, base: Path, index: int, instruction: str, insert: bool) -> dict:
    """Ask the model for one slide: a rewrite of slide `index` (0-based), or (insert=True) a new
    slide to go after it. Returns the slide; the caller puts it into the deck."""
    slides = spec.get("slides", [])
    outline = "\n".join(f"{i + 1}. [{s.get('type')}] {s.get('title', '')}" for i, s in enumerate(slides))
    if insert:
        task = f"Write a NEW slide to insert after slide {index + 1}; do not repeat what other slides say."
    else:
        task = (f"Rewrite slide {index + 1}. Its current content:\n" +
                json.dumps({k: v for k, v in slides[index].items() if not k.startswith("_")}, ensure_ascii=False))
    md = base / "extracted.md"
    content = md.read_text(encoding="utf-8")[:source_limit(model)] if md.exists() else "(not available)"
    prompt = SLIDE_PROMPT.format(schema=SCHEMA, title=spec.get("title", ""), outline=outline, task=task,
                                 instruction=f"INSTRUCTION FROM THE USER: {instruction}\n" if instruction else "",
                                 content=content)
    errors = []
    for attempt in range(3):
        checkpoint()
        log.info("asking %s for one slide (attempt %d)", model, attempt + 1)
        try:
            reply = json.loads(ask_llm(model, prompt))
        except json.JSONDecodeError as e:
            reply, errors = None, [f"reply was not valid JSON: {e}"]
        if isinstance(reply, dict):  # unwrap {"slide": {...}} / {"slides": [{...}]}
            inner = reply.get("slide") or (reply.get("slides") or [None])[0]
            slide = inner if isinstance(inner, dict) and "type" not in reply else reply
            errors = validate({"slides": [slide]}, base)
            if not errors:
                return slide
        prompt += ("\n\nYour previous reply had these problems:\n- " + "\n- ".join(errors) +
                   "\nReply again with the full corrected slide JSON.")
    sys.exit(f"{model} could not write a usable slide: {'; '.join(errors)}")


def count_rule(slides: int, mode: str) -> str:
    """mode: exact = about N, min = at least N (more if the source needs it), auto = model decides."""
    if mode == "auto" or slides <= 0:
        return ("Choose the number of slides the source needs (usually 6-20): cover every important "
                "point, but do not pad.")
    if mode == "min":
        return f"At least {slides} slides - more if the source has enough material, never fewer."
    return f"About {slides} slides."


def auto(files: list[Path], out: Path | None, model: str, slides: int, extra: str,
         preview: bool, animate: bool = True, vision: str | None = "auto",
         workdir: Path | None = None, design: str | None = None, slide_mode: str = "exact") -> Path:
    if is_online(model):
        log.warning("%s is an online model - the extracted document text will be sent to it "
                    "(file reading, OCR and image descriptions still run locally)", model_label(model))
    first = files[0]
    work = workdir or first.parent / f"{first.stem}_ppt"  # extract, assets, spec, previews
    md = extract_all(files, work, vision)
    content = md.read_text(encoding="utf-8")
    if len(content) > (limit := source_limit(model)):
        log.warning("source is %d chars; only the first %d fit in the model's memory window", len(content), limit)
        content = content[:limit]
    prompt = PROMPT.format(schema=SCHEMA, count=count_rule(slides, slide_mode), content=content,
                           extra=f"\nHINTS FROM THE USER (follow them; they override the rules above):\n{extra}\n"
                           if extra else "")
    spec_path = work / "spec.json"
    usable = None  # last spec that renders, even if content_checks still complain about it
    for attempt in range(3):
        log.info("asking %s for a slide spec (attempt %d) - this can take a few minutes", model, attempt + 1)
        try:
            reply = ask_llm(model, prompt)
        except Skipped:
            log.info("skipped further attempts")
            break
        try:
            spec = json.loads(reply)
            errors = validate(spec, work)
        except json.JSONDecodeError as e:
            spec, errors = None, [f"reply was not valid JSON: {e}" +
                                  (" - it is empty or cut off; write fewer, shorter slides and tables of at most 8 rows"
                                   if e.pos >= len(reply.rstrip()) - 2 else "")]
        if spec is not None and not errors:
            usable = spec
            errors = content_checks(spec, content)  # quality complaints: worth a retry, not a failure
            if slide_mode == "min" and slides > 0 and len(spec["slides"]) < slides:
                errors.append(f"only {len(spec['slides'])} slides; write at least {slides} slides")
        if spec is not None:
            spec_path.write_text(json.dumps(spec, indent=2, ensure_ascii=False), encoding="utf-8")
        if not errors or "write" in SKIP:
            break
        log.warning("spec problems: %s", "; ".join(errors))
        prompt += ("\n\nYour previous reply had these problems:\n- " + "\n- ".join(errors) +
                   "\nReply again with the full corrected JSON.")
    if usable is None:
        sys.exit(f"{model} did not write a usable slide spec in {attempt + 1} attempt(s). Try another model "
                 f"(non-thinking models such as qwen2.5 or phi4 are the most reliable), fewer slides, or "
                 f"fix the last attempt by hand ({spec_path}) and run: make_ppt.py build \"{spec_path}\"")
    spec = usable
    slides = spec["slides"]  # drop section dividers with nothing after them
    spec["slides"] = [s for i, s in enumerate(slides) if not (
        s.get("type") == "section" and
        (i == len(slides) - 1 or slides[i + 1].get("type") in ("section", "closing")))]
    if added := add_missing_images(spec, files, content):
        log.info("added slides for images the model skipped: %s", ", ".join(added))
    spec_path.write_text(json.dumps(spec, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Spec: {spec_path} (edit it and run `build` to tweak the deck)")
    online_parts = [w for w, on in (("document text", is_online(model)),
                                    ("pictures", any(is_online(m) for m in VISION_USED))) if on]
    where = (f"ONLINE ({' and '.join(online_parts)} sent out)" if online_parts else "OFFLINE")
    # PowerPoint's Comments property holds at most 255 characters: keep this compact
    spec["made_with"] = " | ".join(filter(None, [
        f"Made {where} by make_ppt.py {__import__('datetime').date.today()}",
        f"Slides: {model_label(model)}",
        f"Images: {', '.join(model_label(m) for m in sorted(VISION_USED)) or 'none'}",
        f"OCR: {', '.join(sorted(OCR_USED)) or 'none'}",
        f"Design: {FAMILIES[design]['name']}" if design in FAMILIES else ""]))
    apply_design(spec, design)
    spec_path.write_text(json.dumps(spec, indent=2, ensure_ascii=False), encoding="utf-8")
    return build(spec_path, out or first.parent / f"{first.stem}.pptx", preview, animate)


# ---------------------------------------------------------------- CLI

def render_family_thumbnails() -> None:
    """Build examples/spec.json in each family and save its cover + a content slide as a
    thumbnail for the app's Design row (ui/families/<key>.jpg)."""
    here = Path(__file__).resolve().parent
    out = here / "ui" / "families"
    out.mkdir(parents=True, exist_ok=True)
    for key in FAMILIES:
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            shutil.copytree(here / "examples" / "assets", work / "assets")
            spec = json.loads((here / "examples" / "spec.json").read_text(encoding="utf-8"))
            spec["family"] = key
            (work / "spec.json").write_text(json.dumps(spec), encoding="utf-8")
            deck = Deck(spec, work)
            deck.build(work / "deck.pptx", animate=False)
            pngs = render_previews(work / "deck.pptx", work / "preview")
            cover = Image.open(pngs[0]).convert("RGB")
            inner = Image.open(pngs[min(7, len(pngs) - 1)]).convert("RGB")  # the steps slide
            w, h = 480, 270
            thumb = cover.resize((w, h))
            thumb.paste(inner.resize((w * 2 // 5, h * 2 // 5)), (w - w * 2 // 5 - 8, h - h * 2 // 5 - 8))  # corner inset
            thumb.save(out / f"{key}.jpg", quality=86)
        print(f"{key}: {out / (key + '.jpg')}")


Deck.lead = staticmethod(lead)  # used by the family layouts (families.py)
Deck.split_even = staticmethod(split_even)


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("extract", help="dump documents to Markdown + images")
    e.add_argument("files", nargs="+", type=Path)
    e.add_argument("-o", "--outdir", type=Path, help="default: <first file>_ppt/ next to it")
    e.add_argument("--vision", default="auto", help="vision model for describing images (default: best installed)")
    e.add_argument("--no-vision", action="store_true", help="do not describe images")
    b = sub.add_parser("build", help="render a JSON spec to .pptx")
    b.add_argument("spec", type=Path)
    b.add_argument("-o", "--out", type=Path)
    b.add_argument("--no-preview", action="store_true")
    b.add_argument("--no-animate", action="store_true", help="no fade transitions/animations")
    b.add_argument("--design", choices=list(FAMILIES), help="design family (see DESIGN_NOTES.md)")
    a = sub.add_parser("auto", help="documents -> local LLM -> .pptx")
    a.add_argument("files", nargs="+", type=Path)
    a.add_argument("-o", "--out", type=Path)
    a.add_argument("-m", "--model", default=DEFAULT_MODEL)
    a.add_argument("--vision", default="auto", help="vision model for describing images (default: best installed)")
    a.add_argument("--no-vision", action="store_true", help="do not describe images")
    a.add_argument("-n", "--slides", type=int, default=10, help="0 = let the model decide")
    a.add_argument("--at-least", action="store_true", help="-n is a minimum: the model may write more")
    a.add_argument("-i", "--instructions", default="", help='e.g. "audience: management, focus on costs"')
    a.add_argument("--no-preview", action="store_true")
    a.add_argument("--no-animate", action="store_true", help="no fade transitions/animations")
    a.add_argument("--design", choices=list(FAMILIES), help="design family (see DESIGN_NOTES.md)")
    sub.add_parser("schema", help="print the slide spec format")
    sub.add_parser("families", help="render a preview of every design family (ui/families/*.jpg)")
    args = ap.parse_args()

    if args.cmd == "extract":
        md = extract_all(args.files, args.outdir or args.files[0].parent / f"{args.files[0].stem}_ppt",
                         None if args.no_vision else args.vision)
        print(f"Extract: {md}")
    elif args.cmd == "build":
        build(args.spec, args.out, not args.no_preview, not args.no_animate, args.design)
    elif args.cmd == "auto":
        auto(args.files, args.out, args.model, args.slides, args.instructions, not args.no_preview,
             not args.no_animate, None if args.no_vision else args.vision, design=args.design,
             slide_mode="min" if args.at_least else "exact")
    elif args.cmd == "families":
        render_family_thumbnails()
    else:
        print(SCHEMA)


if __name__ == "__main__":
    main()

