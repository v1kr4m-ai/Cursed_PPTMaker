"""Design families: original slide compositions (see DESIGN_NOTES.md).

Each family is a palette, a font pairing, a few style flags, and a choice of layout
variants. The variants themselves are drawn by FamilyLayouts (mixed into make_ppt.Deck),
using plain PowerPoint shapes so decks stay editable and fully offline.
"""
from __future__ import annotations

import math

from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FAMILIES = {
    "soft": {
        "name": "Soft Dashboard",
        "about": "Raised white cards on soft grey, ring gauges, dot header with page number",
        "bg": "EEF1F5", "card": "FFFFFF", "ink": "1F2430", "muted": "6B7385", "dark": "1F2430",
        "accents": ["F2557E", "F9A43C", "2EC4A6", "3D7BFA"],
        "head": "Century Gothic", "body": "Segoe UI", "num": "Century Gothic",
        "soft": True, "caps": False, "radius": 0.1,
        "header": "dots_page", "title": "soft", "section": "soft", "steps": "rows",
        "cards": "cards", "stats": "rings", "closing": "light",
    },
    "ribbon": {
        "name": "Ribbon Steps",
        "about": "Banner ribbons with 01 STEP numerals, diamond cards, condensed capitals",
        "bg": "FFFFFF", "card": "F4F5F7", "ink": "22252B", "muted": "6E7480", "dark": "1E2230",
        "accents": ["22A45D", "F28C1E", "2D8FE0", "E5484D"],
        "head": "Bahnschrift SemiBold Condensed", "body": "Segoe UI", "num": "Bahnschrift SemiBold Condensed",
        "soft": False, "caps": True, "radius": 0.06,
        "header": "stripes", "title": "ribbon", "section": "bignum", "steps": "ribbons",
        "cards": "diamonds", "stats": "cards", "closing": "dark",
    },
    "hexa": {
        "name": "Hexa Hub",
        "about": "Hexagon cards with coloured edges, segmented ring cover, giant section numbers",
        "bg": "E9EEF2", "card": "FFFFFF", "ink": "263238", "muted": "66737D", "dark": "22303A",
        "accents": ["F2B01E", "8E3B8F", "4CB050", "2E6FD1"],
        "head": "Bahnschrift SemiBold Condensed", "body": "Segoe UI", "num": "Bahnschrift SemiBold Condensed",
        "soft": True, "caps": True, "radius": 0.08,
        "header": "hexdots", "title": "hub", "section": "bignum", "steps": "hexsteps",
        "cards": "hexagons", "stats": "rings", "closing": "dark",
    },
    "editorial": {
        "name": "Editorial Mono",
        "about": "Black and white with a diagonal split cover, serif headlines, thin rules",
        "bg": "FFFFFF", "card": "F3F2EF", "ink": "141414", "muted": "6A6A6A", "dark": "141414",
        "accents": ["4B7FA3", "141414", "8F8F8F", "B89A5E"],
        "head": "Georgia", "body": "Segoe UI", "num": "Georgia",
        "soft": False, "caps": False, "radius": 0.0,
        "header": "rule", "title": "split", "section": "split", "steps": "rows",
        "cards": "cards", "stats": "cards", "closing": "dark",
    },
    "wander": {
        "name": "Wanderlust",
        "about": "Warm paper, serif headlines, organic circles and teal round badges",
        "bg": "F5F1EA", "card": "FFFFFF", "ink": "1E2B2F", "muted": "6C7A7D", "dark": "173A40",
        "accents": ["2B8C95", "C98B4E", "175A63", "E07A5F"],
        "head": "Georgia", "body": "Segoe UI", "num": "Georgia",
        "soft": True, "caps": False, "radius": 0.14,
        "header": "kicker", "title": "wander", "section": "wander", "steps": "rows",
        "cards": "circles", "stats": "cards", "closing": "light",
    },
    "flow": {
        "name": "Timeline Flow",
        "about": "Chevron arrow steps, dark band cover, accent-bar headers",
        "bg": "F7F9FC", "card": "FFFFFF", "ink": "1C2B3A", "muted": "5E6E80", "dark": "14324A",
        "accents": ["1E88C8", "22B3A6", "0F4C75", "F5A623"],
        "head": "Segoe UI Semibold", "body": "Segoe UI", "num": "Bahnschrift SemiBold Condensed",
        "soft": False, "caps": False, "radius": 0.06,
        "header": "kicker_bar", "title": "band", "section": "bignum", "steps": "chevrons",
        "cards": "cards", "stats": "cards", "closing": "dark",
    },
}


def names() -> dict[str, str]:
    return {k: v["name"] for k, v in FAMILIES.items()}


# ---------------------------------------------------------------- layouts


W, H = 13.333, 7.5
LEFT, RIGHT, TOP, BOTTOM = 0.6, 12.733, 1.65, 6.9
CW = RIGHT - LEFT
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def shade(c, amount):
    """Darken a colour towards black."""
    return RGBColor(*(round(v * (1 - amount)) for v in c))


def tint(c, amount):
    """Lighten a colour towards white."""
    return RGBColor(*(round(v + (255 - v) * amount) for v in c))


class FamilyLayouts:
    """Layout variants used by the families. Mixed into make_ppt.Deck, so it uses the deck's
    primitives (text, box, badge, mark, new_slide, header) and colours (self.ACCENTS etc.)."""

    # helpers -------------------------------------------------------------
    def shape(self, slide, kind, x, y, w, h, fill, rot=0, raised=False, line=None, line_w=1.5):
        s = self.box(slide, x, y, w, h, fill, kind, raised=raised)
        if rot:
            s.rotation = rot
        if line is not None:
            s.line.color.rgb = line
            s.line.width = Pt(line_w)
        return s

    def arc(self, slide, cx, cy, d, start, end, thick, fill):
        """Ring segment from `start` to `end` degrees, clockwise (0 = 3 o'clock, 270 = top)."""
        s = self.box(slide, cx - d / 2, cy - d / 2, d, d, fill, MSO_SHAPE.BLOCK_ARC, raised=False)
        s.adjustments[0], s.adjustments[1], s.adjustments[2] = (start % 360) * 0.6, (end % 360) * 0.6, thick
        return s

    def ring(self, slide, cx, cy, d, thick, fill):
        s = self.box(slide, cx - d / 2, cy - d / 2, d, d, fill, MSO_SHAPE.DONUT, raised=False)
        s.adjustments[0] = thick
        return s

    def cap(self, text):
        return str(text).upper() if self.caps else str(text)

    def acc(self, i):
        return self.ACCENTS[i % len(self.ACCENTS)]

    def centered(self, slide, x, y, w, h, text, size, color, font=None, bold=True):
        return self.text(slide, x, y, w, h, str(text), size=size, color=color, bold=bold, font=font,
                         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, min_size=max(8, size * 0.5))

    def footer(self, slide, s, x=0.8, color=None):
        if s.get("footer"):
            self.text(slide, x, 6.6, 11, 0.4, s["footer"], size=13, color=color or self.MUTED)

    def first_image(self):
        for sl in self.spec.get("slides", []):
            if isinstance(sl.get("image"), str) and (self.base / sl["image"]).exists():
                return sl["image"]
        return None

    def picture_circle(self, slide, path, x, y, d):
        """Picture cropped to a square and shown in a circle."""
        path = self.base / path
        with Image.open(path) as im:
            iw, ih = im.size
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(d), Inches(d))
        if iw > ih:
            pic.crop_left = pic.crop_right = (1 - ih / iw) / 2
        else:
            pic.crop_top = pic.crop_bottom = (1 - iw / ih) / 2
        pic._element.spPr.find(qn("a:prstGeom")).set("prst", "ellipse")
        return pic

    # covers --------------------------------------------------------------
    def t_soft(self, s):
        slide = self.new_slide(s)
        for i in range(3):
            self.box(slide, 0.9 + i * 0.32, 1.55, 0.2, 0.2, self.acc(i), MSO_SHAPE.OVAL, raised=False)
        self.text(slide, 0.9, 1.9, 6.9, 2.3, s["title"], size=50, color=self.HEAD, bold=True, font=self.HFONT,
                  anchor=MSO_ANCHOR.BOTTOM, min_size=30)
        if s.get("subtitle"):
            self.text(slide, 0.9, 4.4, 7.0, 1.3, s["subtitle"], size=21, color=self.MUTED, min_size=13)
        self.footer(slide, s, 0.9)
        self.mark(slide)
        cx, cy = 10.5, 3.5
        self.box(slide, cx - 2.0, cy - 2.0, 4.0, 4.0, self.TINT, MSO_SHAPE.OVAL)
        self.ring(slide, cx, cy, 3.0, 0.09, tint(self.MUTED, 0.75))
        self.arc(slide, cx, cy, 3.0, 270, 200, 0.09, self.acc(0))
        self.box(slide, cx - 0.55, cy - 0.55, 1.1, 1.1, self.acc(3), MSO_SHAPE.OVAL, raised=False)
        self.box(slide, 10.9, 5.75, 1.9, 0.8, self.TINT, radius=0.3)
        self.box(slide, 11.15, 6.02, 1.2, 0.08, self.acc(1), MSO_SHAPE.RECTANGLE, raised=False)
        self.box(slide, 11.15, 6.2, 0.8, 0.08, tint(self.MUTED, 0.6), MSO_SHAPE.RECTANGLE, raised=False)

    def t_ribbon(self, s):
        slide = self.new_slide(s)
        self.text(slide, 0.8, 1.5, 6.9, 2.7, self.cap(s["title"]), size=60, color=self.HEAD, bold=True,
                  font=self.HFONT, anchor=MSO_ANCHOR.BOTTOM, min_size=32)
        for i in range(4):
            self.box(slide, 0.8 + i * 0.5, 4.38, 0.4, 0.09, self.acc(i), MSO_SHAPE.RECTANGLE, raised=False)
        if s.get("subtitle"):
            self.text(slide, 0.8, 4.7, 6.6, 1.4, s["subtitle"], size=21, color=self.MUTED, min_size=13)
        self.footer(slide, s)
        for i, h in enumerate((4.9, 4.0, 4.5, 3.6)):
            self.mark(slide)
            x = 8.0 + i * 1.22
            self.box(slide, x, 0, 1.0, h, self.acc(i), MSO_SHAPE.FLOWCHART_OFFPAGE_CONNECTOR, raised=False)
            self.centered(slide, x, h - 1.55, 1.0, 0.7, f"{i + 1:02d}", 30, _WHITE, self.NFONT)
            self.centered(slide, x, h - 0.95, 1.0, 0.3, "STEP", 11, _WHITE)

    def t_hub(self, s):
        slide = self.new_slide(s)
        self.text(slide, 0.8, 1.7, 6.2, 2.4, self.cap(s["title"]), size=56, color=self.HEAD, bold=True,
                  font=self.HFONT, anchor=MSO_ANCHOR.BOTTOM, min_size=30)
        for i in range(5):
            self.box(slide, 0.8 + i * 0.34, 4.3, 0.2, 0.18, self.acc(i), MSO_SHAPE.HEXAGON, raised=False)
        if s.get("subtitle"):
            self.text(slide, 0.8, 4.7, 6.0, 1.4, s["subtitle"], size=20, color=self.MUTED, min_size=13)
        self.footer(slide, s)
        self.mark(slide)
        cx, cy = 9.7, 3.75
        for i in range(4):
            start = 270 + i * 90 + 6
            self.arc(slide, cx, cy, 4.7, start, start + 78, 0.12, self.acc(i))
        self.box(slide, cx - 1.6, cy - 1.6, 3.2, 3.2, self.TINT, MSO_SHAPE.OVAL)
        self.centered(slide, cx - 1.2, cy - 0.9, 2.4, 1.8, (s["title"].strip() or "·")[0].upper(), 80,
                      self.INK, self.NFONT)

    def t_split(self, s):
        slide = self.new_slide(s)
        self.box(slide, 0, 0, 6.2, H, self.INK, MSO_SHAPE.RECTANGLE, raised=False)
        self.box(slide, 6.19, 0, 1.6, H, self.INK, MSO_SHAPE.RIGHT_TRIANGLE, raised=False)  # diagonal edge
        self.text(slide, 0.8, 1.6, 5.2, 2.9, s["title"], size=50, color=_WHITE, bold=True, font=self.HFONT,
                  anchor=MSO_ANCHOR.BOTTOM, min_size=28)
        self.box(slide, 0.8, 4.75, 0.6, 0.08, self.ACCENT, MSO_SHAPE.RECTANGLE, raised=False)
        self.mark(slide)
        self.box(slide, 8.4, 2.85, 1.2, 0.03, self.ACCENT, MSO_SHAPE.RECTANGLE, raised=False)
        if s.get("subtitle"):
            self.text(slide, 8.4, 3.05, 4.3, 2.0, s["subtitle"], size=22, color=self.HEAD, font=self.HFONT,
                      min_size=13)
        self.footer(slide, s, 8.4)

    def t_wander(self, s):
        slide = self.new_slide(s)
        if s.get("kicker"):
            self.text(slide, 0.8, 1.55, 6.2, 0.35, s["kicker"].upper(), size=12, color=self.ACCENT, bold=True)
        self.text(slide, 0.8, 1.9, 6.6, 2.4, s["title"], size=54, color=self.HEAD, bold=True, font=self.HFONT,
                  anchor=MSO_ANCHOR.BOTTOM, min_size=30)
        self.box(slide, 0.8, 4.45, 0.9, 0.05, self.ACCENT, MSO_SHAPE.RECTANGLE, raised=False)
        if s.get("subtitle"):
            self.text(slide, 0.8, 4.7, 6.2, 1.3, s["subtitle"], size=20, color=self.MUTED, min_size=13)
        self.footer(slide, s)
        self.mark(slide)
        self.box(slide, 7.4, 4.7, 2.2, 2.2, tint(self.ACCENT2, 0.35), MSO_SHAPE.OVAL, raised=False)
        self.box(slide, 8.2, 0.9, 5.6, 5.6, self.ACCENT, MSO_SHAPE.OVAL, raised=False)
        self.box(slide, 12.2, 0.45, 0.8, 0.8, self.ACCENT3, MSO_SHAPE.OVAL, raised=False)
        if img := self.first_image():
            self.picture_circle(slide, img, 8.5, 1.2, 5.0)

    def t_band(self, s):
        slide = self.new_slide(s)
        self.box(slide, 0, 0, W, 4.7, self.INK, MSO_SHAPE.RECTANGLE, raised=False)
        self.box(slide, 0, 4.7, W, 0.12, self.ACCENT, MSO_SHAPE.RECTANGLE, raised=False)
        self.text(slide, 0.8, 1.0, 10.5, 2.6, s["title"], size=52, color=_WHITE, bold=True, font=self.HFONT,
                  anchor=MSO_ANCHOR.BOTTOM, min_size=30)
        if s.get("subtitle"):
            self.text(slide, 0.8, 3.7, 10.5, 0.8, s["subtitle"], size=20, color=tint(self.INK, 0.7), min_size=13)
        self.mark(slide)
        for i in range(4):
            kind = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
            self.box(slide, 0.8 + i * 1.5, 5.35, 1.75, 0.75, self.acc(i), kind, raised=False)
        self.footer(slide, s, 7.6)

    # section dividers ------------------------------------------------------
    def sec_bignum(self, s):
        slide = self.new_slide(s)
        self.section_no += 1
        self.box(slide, 0, 0, 4.8, H, self.INK, MSO_SHAPE.RECTANGLE, raised=False)
        self.centered(slide, 0.3, 1.6, 4.2, 4.0, f"{self.section_no:02d}", 160, self.ACCENT, self.NFONT)
        self.mark(slide)
        self.text(slide, 5.5, 2.2, 7.2, 2.0, self.cap(s["title"]), size=46, color=self.HEAD, bold=True,
                  font=self.HFONT, anchor=MSO_ANCHOR.BOTTOM, min_size=28)
        self.box(slide, 5.5, 4.35, 1.1, 0.08, self.ACCENT2, MSO_SHAPE.RECTANGLE, raised=False)
        if s.get("subtitle"):
            self.text(slide, 5.5, 4.6, 7.0, 1.4, s["subtitle"], size=20, color=self.MUTED, min_size=13)

    def sec_soft(self, s):
        slide = self.new_slide(s)
        self.section_no += 1
        self.box(slide, 2.2, 1.7, 8.9, 4.1, self.TINT, radius=0.08)
        self.badge(slide, 2.8, 2.3, s.get("icon") or self.section_no, self.ACCENT, d=1.0, size=26)
        self.text(slide, 2.8, 3.55, 7.8, 1.1, s["title"], size=40, color=self.HEAD, bold=True, font=self.HFONT,
                  min_size=24)
        if s.get("subtitle"):
            self.text(slide, 2.8, 4.65, 7.8, 0.9, s["subtitle"], size=19, color=self.MUTED, min_size=12)
        for i in range(3):
            self.box(slide, 10.0 + i * 0.3, 2.45, 0.17, 0.17, self.acc(i), MSO_SHAPE.OVAL, raised=False)

    def sec_split(self, s):
        slide = self.new_slide(s)
        self.section_no += 1
        self.box(slide, 0, 0, 5.0, H, self.INK, MSO_SHAPE.RECTANGLE, raised=False)
        self.box(slide, 4.99, 0, 1.2, H, self.INK, MSO_SHAPE.RIGHT_TRIANGLE, raised=False)  # diagonal edge
        self.text(slide, 0.8, 2.0, 4.0, 2.4, f"{self.section_no:02d}", size=120, color=_WHITE, font=self.NFONT,
                  anchor=MSO_ANCHOR.BOTTOM, min_size=60)
        self.mark(slide)
        self.text(slide, 6.8, 2.4, 5.9, 1.9, s["title"], size=44, color=self.HEAD, bold=True, font=self.HFONT,
                  anchor=MSO_ANCHOR.BOTTOM, min_size=26)
        self.box(slide, 6.8, 4.45, 1.2, 0.03, self.ACCENT, MSO_SHAPE.RECTANGLE, raised=False)
        if s.get("subtitle"):
            self.text(slide, 6.8, 4.65, 5.9, 1.4, s["subtitle"], size=19, color=self.MUTED, min_size=12)

    def sec_wander(self, s):
        slide = self.new_slide(s)
        self.section_no += 1
        self.box(slide, 9.3, 0.8, 5.6, 5.6, tint(self.ACCENT, 0.15), MSO_SHAPE.OVAL, raised=False)
        self.box(slide, 8.6, 5.2, 1.6, 1.6, tint(self.ACCENT2, 0.3), MSO_SHAPE.OVAL, raised=False)
        self.text(slide, 0.8, 2.15, 7.0, 0.4, f"CHAPTER {self.section_no:02d}", size=13, color=self.ACCENT, bold=True)
        self.text(slide, 0.8, 2.55, 7.6, 1.7, s["title"], size=48, color=self.HEAD, bold=True, font=self.HFONT,
                  min_size=28)
        self.box(slide, 0.8, 4.35, 0.9, 0.05, self.ACCENT, MSO_SHAPE.RECTANGLE, raised=False)
        if s.get("subtitle"):
            self.text(slide, 0.8, 4.6, 7.2, 1.3, s["subtitle"], size=19, color=self.MUTED, min_size=12)

    # steps -------------------------------------------------------------------
    def st_ribbons(self, s):
        num = 0
        for n, chunk in enumerate(self.split_even(s["steps"], 5)):
            slide = self.header(s, s["title"] + (" (cont.)" if n else ""))
            k, gap, bh = len(chunk), 0.3, 1.75
            cw = (CW - gap * (k - 1)) / k
            for i, st in enumerate(chunk):
                self.mark(slide)
                num += 1
                x, col = LEFT + i * (cw + gap), self.acc(num - 1)
                self.box(slide, x, TOP - 0.05, cw, bh, col, MSO_SHAPE.FLOWCHART_OFFPAGE_CONNECTOR, raised=False)
                self.centered(slide, x, TOP + 0.05, cw, 0.8, f"{num:02d}", 40, _WHITE, self.NFONT)
                self.centered(slide, x, TOP + 0.85, cw, 0.3, "STEP", 12, _WHITE)
                y = TOP + bh + 0.2
                self.box(slide, x, y, cw, BOTTOM - y, self.TINT)
                self.text(slide, x + 0.2, y + 0.2, cw - 0.4, 0.65, self.cap(st.get("title", "")), size=18,
                          color=self.HEAD, bold=True, font=self.HFONT, min_size=12)
                self.text(slide, x + 0.2, y + 0.9, cw - 0.4, BOTTOM - y - 1.05, st.get("text", ""), size=14,
                          min_size=10)

    def st_chevrons(self, s):
        num = 0
        for n, chunk in enumerate(self.split_even(s["steps"], 5)):
            slide = self.header(s, s["title"] + (" (cont.)" if n else ""))
            k, overlap, ch = len(chunk), 0.3, 1.05
            w = (CW + overlap * (k - 1)) / k
            for i, st in enumerate(chunk):
                self.mark(slide)
                num += 1
                x = LEFT + i * (w - overlap)
                kind = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
                self.box(slide, x, TOP + 0.1, w, ch, self.acc(num - 1), kind, raised=False)
                self.centered(slide, x + 0.15, TOP + 0.1, w - 0.3, ch, f"{num:02d}", 30, _WHITE, self.NFONT)
                y = TOP + ch + 0.45
                self.text(slide, x + 0.3, y, w - 0.6, 0.65, st.get("title", ""), size=18, color=self.HEAD,
                          bold=True, font=self.HFONT, min_size=12)
                self.text(slide, x + 0.3, y + 0.7, w - 0.6, BOTTOM - y - 0.7, st.get("text", ""), size=14,
                          min_size=10)

    def st_hexsteps(self, s):
        num = 0
        for n, chunk in enumerate(self.split_even(s["steps"], 5)):
            slide = self.header(s, s["title"] + (" (cont.)" if n else ""))
            k, gap, hw, hh = len(chunk), 0.3, 1.8, 1.56
            cw = (CW - gap * (k - 1)) / k
            for i, st in enumerate(chunk):
                self.mark(slide)
                num += 1
                x = LEFT + i * (cw + gap)
                hx, hy = x + (cw - hw) / 2, TOP + 0.1
                self.box(slide, hx - 0.1, hy - 0.09, hw + 0.2, hh + 0.18, self.acc(num - 1), MSO_SHAPE.HEXAGON,
                         raised=False)
                self.box(slide, hx, hy, hw, hh, self.TINT, MSO_SHAPE.HEXAGON)
                self.centered(slide, hx, hy, hw, hh, f"{num:02d}", 34, self.acc(num - 1), self.NFONT)
                y = hy + hh + 0.35
                self.text(slide, x, y, cw, 0.6, self.cap(st.get("title", "")), size=18, color=self.HEAD, bold=True,
                          font=self.HFONT, align=PP_ALIGN.CENTER, min_size=12)
                self.text(slide, x + 0.1, y + 0.65, cw - 0.2, BOTTOM - y - 0.65, st.get("text", ""), size=14,
                          align=PP_ALIGN.CENTER, min_size=10)

    # cards -------------------------------------------------------------------
    def cd_diamonds(self, s):
        cards = s["cards"]
        if len(cards) > 4:
            return self.cards_default(s)
        slide = self.header(s)
        k, gap, d = len(cards), 0.3, 2.0
        cw = (CW - gap * (k - 1)) / k
        for i, c in enumerate(cards):
            self.mark(slide)
            x = LEFT + i * (cw + gap)
            dx, dy = x + (cw - d) / 2, TOP + 0.15 + (0.4 if i % 2 else 0)
            self.box(slide, dx + 0.14, dy + 0.14, d, d, self.acc(i), MSO_SHAPE.DIAMOND, raised=False)
            self.shape(slide, MSO_SHAPE.DIAMOND, dx, dy, d, d, _WHITE, line=tint(self.MUTED, 0.6), line_w=0.75)
            self.centered(slide, dx, dy, d, d, c.get("icon") or f"{i + 1:02d}", 32, self.acc(i), self.NFONT)
            y = dy + d + 0.3
            self.text(slide, x, y, cw, 0.6, self.cap(c.get("title", "")), size=19, color=self.HEAD, bold=True,
                      font=self.HFONT, align=PP_ALIGN.CENTER, min_size=12)
            self.text(slide, x + 0.1, y + 0.65, cw - 0.2, BOTTOM - y - 0.65, c.get("text", ""), size=14,
                      align=PP_ALIGN.CENTER, min_size=10)

    def cd_hexagons(self, s):
        cards = s["cards"][:6]
        slide = self.header(s)
        if len(cards) <= 4:  # one zig-zag row, text under each hexagon
            k, gap, hw, hh = len(cards), 0.3, 2.0, 1.74
            cw = (CW - gap * (k - 1)) / k
            for i, c in enumerate(cards):
                self.mark(slide)
                x = LEFT + i * (cw + gap)
                hx, hy = x + (cw - hw) / 2, TOP + 0.1 + (0.35 if i % 2 else 0)
                self.box(slide, hx - 0.1, hy - 0.09, hw + 0.2, hh + 0.18, self.acc(i), MSO_SHAPE.HEXAGON, raised=False)
                self.box(slide, hx, hy, hw, hh, self.TINT, MSO_SHAPE.HEXAGON)
                self.centered(slide, hx, hy, hw, hh, c.get("icon") or f"{i + 1:02d}", 34, self.acc(i), self.NFONT)
                y = hy + hh + 0.3
                self.text(slide, x, y, cw, 0.6, self.cap(c.get("title", "")), size=19, color=self.HEAD, bold=True,
                          font=self.HFONT, align=PP_ALIGN.CENTER, min_size=12)
                self.text(slide, x + 0.1, y + 0.65, cw - 0.2, BOTTOM - y - 0.65, c.get("text", ""), size=14,
                          align=PP_ALIGN.CENTER, min_size=10)
        else:  # 2 rows x 3: hexagon on the left of its text
            gap, hw, hh = 0.3, 1.25, 1.09
            cw, rh = (CW - gap * 2) / 3, (BOTTOM - TOP - gap) / 2
            for i, c in enumerate(cards):
                self.mark(slide)
                x, y = LEFT + (i % 3) * (cw + gap), TOP + (i // 3) * (rh + gap)
                self.box(slide, x - 0.07, y - 0.06, hw + 0.14, hh + 0.12, self.acc(i), MSO_SHAPE.HEXAGON, raised=False)
                self.box(slide, x, y, hw, hh, self.TINT, MSO_SHAPE.HEXAGON)
                self.centered(slide, x, y, hw, hh, c.get("icon") or f"{i + 1:02d}", 24, self.acc(i), self.NFONT)
                self.text(slide, x + hw + 0.25, y, cw - hw - 0.3, 0.6, self.cap(c.get("title", "")), size=17,
                          color=self.HEAD, bold=True, font=self.HFONT, min_size=11)
                self.text(slide, x + hw + 0.25, y + 0.6, cw - hw - 0.3, rh - 0.6, c.get("text", ""), size=13,
                          min_size=9)

    def cd_circles(self, s):
        cards = s["cards"][:6]
        slide = self.header(s)
        cols = 2 if len(cards) == 4 else min(len(cards), 3)
        rows = math.ceil(len(cards) / cols)
        gap = 0.45
        cw, rh = (CW - gap * (cols - 1)) / cols, (BOTTOM - TOP - gap * (rows - 1)) / rows
        for i, c in enumerate(cards):
            self.mark(slide)
            x, y = LEFT + (i % cols) * (cw + gap), TOP + (i // cols) * (rh + gap)
            d = 0.95 if rows == 1 else 0.75
            self.box(slide, x, y, d, d, self.acc(i), MSO_SHAPE.OVAL, raised=False)
            self.centered(slide, x, y, d, d, c.get("icon") or i + 1, 24 if rows == 1 else 18, _WHITE, self.NFONT)
            if rows == 1:
                self.text(slide, x, y + d + 0.3, cw, 0.7, c.get("title", ""), size=22, color=self.HEAD, bold=True,
                          font=self.HFONT, min_size=14)
                self.box(slide, x, y + d + 1.05, 0.6, 0.04, self.acc(i), MSO_SHAPE.RECTANGLE, raised=False)
                self.text(slide, x, y + d + 1.25, cw, rh - d - 1.3, c.get("text", ""), size=15, min_size=10)
            else:
                self.text(slide, x + d + 0.25, y, cw - d - 0.3, 0.6, c.get("title", ""), size=19, color=self.HEAD,
                          bold=True, font=self.HFONT, min_size=12)
                self.text(slide, x + d + 0.25, y + 0.6, cw - d - 0.3, rh - 0.6, c.get("text", ""), size=14, min_size=9)

    # stats -------------------------------------------------------------------
    def sa_rings(self, s):
        stats = s["stats"][:4]
        slide = self.header(s)
        k, gap, d = len(stats), 0.3, 2.2
        cw = (CW - gap * (k - 1)) / k
        card_h = d + 1.75
        for i, st in enumerate(stats):
            self.mark(slide)
            x = LEFT + i * (cw + gap)
            if self.soft:
                self.box(slide, x, TOP, cw, card_h, self.TINT)
            cx, cy = x + cw / 2, TOP + 0.3 + d / 2
            value = str(st.get("value", ""))
            try:
                p = float(value.replace(",", "").rstrip("%")) / 100 if value.strip().endswith("%") else 0.75
            except ValueError:
                p = 0.75
            p = min(max(p, 0.02), 1.0)
            self.ring(slide, cx, cy, d, 0.1, tint(self.MUTED, 0.78))
            if p >= 0.995:
                self.ring(slide, cx, cy, d, 0.1, self.acc(i))
            else:
                self.arc(slide, cx, cy, d, 270, 270 + 360 * p, 0.1, self.acc(i))
            self.centered(slide, cx - d / 2 + 0.25, cy - 0.5, d - 0.5, 1.0, value, 30, self.HEAD, self.NFONT)
            self.text(slide, x + 0.2, cy + d / 2 + 0.2, cw - 0.4, 0.9, st.get("label", ""), size=15,
                      color=self.MUTED, align=PP_ALIGN.CENTER, min_size=10)
        self.mark(slide)
        if s.get("text"):
            y = TOP + card_h + 0.3
            self.text(slide, LEFT, y, CW, BOTTOM - y, [self.lead(s["text"])], size=17, min_size=11)

    # closing -------------------------------------------------------------------
    def cl_light(self, s):
        slide = self.new_slide(s)
        items = s.get("items", [])[:6]
        self.text(slide, 0.8, 0.7, 11.7, 0.9, s["title"], size=40, color=self.HEAD, bold=True, font=self.HFONT,
                  min_size=26)
        step = min(0.9, 4.9 / max(len(items), 1))
        for i, item in enumerate(items):
            self.mark(slide)
            y = 1.95 + i * step
            if self.soft:
                self.box(slide, 0.8, y - 0.08, 11.7, step - 0.14, self.TINT, radius=0.25)
            self.badge(slide, 1.0, y + (step - 0.14 - 0.55) / 2 - 0.08, "✓", self.acc(i), d=0.55, size=18)
            self.text(slide, 1.85, y - 0.08, 10.4, step - 0.14, item, size=21, color=self.BODY,
                      anchor=MSO_ANCHOR.MIDDLE, min_size=13)
